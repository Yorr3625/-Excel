import asyncio
import json
from io import BytesIO
from types import SimpleNamespace

import pytest
from openpyxl import Workbook, load_workbook
from starlette.applications import Starlette
from starlette.responses import JSONResponse
from starlette.routing import Route
from starlette.testclient import TestClient
from reflex._upload import upload

import orders_dashboard.orders_dashboard as dashboard


def server_state(token):
    return SimpleNamespace(
        is_authenticated=False,
        router=SimpleNamespace(headers=SimpleNamespace(cookie=f"dashboard_session={token}")),
        status="", selected_file="", uploaded_file_path="", order_mail_metadata=None,
        order_record_warning="", output_file="", log_file="", error_text="", duplicate_note="",
        _reset_preview=lambda: None, detect_and_set_mode=lambda: None,
    )


class StateStore:
    def __init__(self, state):
        self.state = state

    async def get_state(self, cls):
        assert cls is dashboard.State
        return self.state


@pytest.mark.parametrize("handler", ["handle_upload", "recognize_invoice_photos"])
def test_upload_event_without_router_uses_live_server_session(handler):
    token = dashboard.dashboard_sessions.create_session()
    try:
        state = server_state(token)
        event = SimpleNamespace(name=f"{dashboard.State.get_full_name()}.{handler}", router_data={})
        result = asyncio.run(dashboard.DashboardAuthMiddleware().preprocess(None, StateStore(state), event))
        assert result is None
        assert state.is_authenticated
    finally:
        dashboard.dashboard_sessions.destroy_session(token)


def test_revoked_session_cannot_authorize_upload_from_cached_state():
    token = dashboard.dashboard_sessions.create_session()
    dashboard.dashboard_sessions.destroy_session(token)
    state = server_state(token)
    state.is_authenticated = True
    event = SimpleNamespace(name=f"{dashboard.State.get_full_name()}.handle_upload", router_data={})
    update = asyncio.run(dashboard.DashboardAuthMiddleware().preprocess(None, StateStore(state), event))
    assert update is not None
    assert update.events[0].name == "_redirect"


def test_normal_event_without_headers_cannot_use_upload_fallback():
    token = dashboard.dashboard_sessions.create_session()
    try:
        state = server_state(token)
        event = SimpleNamespace(name=f"{dashboard.State.get_full_name()}.set_page", router_data={})
        assert asyncio.run(dashboard.DashboardAuthMiddleware().preprocess(None, StateStore(state), event)) is not None
    finally:
        dashboard.dashboard_sessions.destroy_session(token)


def test_http_upload_requires_real_cookie_before_parsing_files():
    calls = []

    async def endpoint(request):
        calls.append(request)
        return JSONResponse({"ok": True})

    client = TestClient(dashboard.DashboardUploadAuth(Starlette(routes=[Route("/_upload", endpoint, methods=["POST"])])))
    assert client.post("/_upload", files={"files": ("order.xlsx", b"synthetic")}).status_code == 401
    client.cookies.set("dashboard_session", "invalid")
    assert client.post("/_upload", files={"files": ("order.xlsx", b"synthetic")}).status_code == 401
    assert not calls
    token = dashboard.dashboard_sessions.create_session()
    try:
        client.cookies.set("dashboard_session", token)
        assert client.post("/_upload", files={"files": ("order.xlsx", b"synthetic")}).status_code == 200
        assert len(calls) == 1
    finally:
        dashboard.dashboard_sessions.destroy_session(token)


def test_real_reflex_multipart_upload_passes_auth_and_saves_excel(tmp_path, monkeypatch):
    monkeypatch.setattr(dashboard.rx, "get_upload_dir", lambda: tmp_path)
    monkeypatch.setattr(dashboard, "was_processed", lambda filename: None)
    token = dashboard.dashboard_sessions.create_session()
    state = server_state(token)
    events = []

    class EventProcessor:
        async def enqueue_stream_delta(self, client_token, event):
            assert client_token == "synthetic-client"
            events.append(event)
            assert event.router_data == {}
            result = await dashboard.DashboardAuthMiddleware().preprocess(None, StateStore(state), event)
            assert result is None, "Authenticated upload was redirected to login"
            await dashboard.State.handle_upload.fn(state, event.payload["files"])
            yield {dashboard.State.get_full_name(): {"selected_file": state.selected_file, "status": state.status}}

    app = SimpleNamespace(event_processor=EventProcessor())
    client = TestClient(dashboard.DashboardUploadAuth(Starlette(routes=[Route("/_upload", upload(app), methods=["POST"])])))
    client.cookies.set("dashboard_session", token)
    book = Workbook()
    book.active.append(["Товар", "Магазин"])
    book.active.append(["Яблоки", 12])
    stream = BytesIO()
    book.save(stream)
    try:
        response = client.post("/_upload", files={"files": ("заказ.xlsx", stream.getvalue())},
                               headers={"reflex-client-token": "synthetic-client",
                                        "reflex-event-handler": f"{dashboard.State.get_full_name()}.handle_upload"})
        assert response.status_code == 200
        assert len(events) == 1
        assert state.is_authenticated
        assert state.selected_file == "заказ.xlsx"
        saved = load_workbook(tmp_path / "заказ.xlsx")
        assert saved.active.cell(2, 2).value == 12
        saved.close()
        payload = json.loads(response.text)
        assert payload["delta"][dashboard.State.get_full_name()]["selected_file"] == "заказ.xlsx"
        assert "_redirect" not in response.text
    finally:
        dashboard.dashboard_sessions.destroy_session(token)
