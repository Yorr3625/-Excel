import asyncio
import json
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import Workbook, load_workbook
from starlette.datastructures import UploadFile

import pytest
from starlette.applications import Starlette
from starlette.routing import Route
from starlette.testclient import TestClient

import orders_dashboard.orders_dashboard as dashboard
from modules import dashboard_sessions, mail_watcher, order_records, paths
from orders_dashboard.api.admin import download_order_source_file


def upload_state():
    return SimpleNamespace(
        status="", selected_file="", uploaded_file_path="", order_mail_metadata={"old": True},
        order_record_warning="", output_file="old.xlsx", log_file="old.log", error_text="old error",
        duplicate_note="", _reset_preview=lambda: None, detect_and_set_mode=lambda: None,
    )


def test_manual_excel_upload_saves_bytes_and_updates_selected_file(tmp_path, monkeypatch):
    book = Workbook()
    book.active.append(["Товар", "Магазин"])
    book.active.append(["Яблоки", 12])
    content = BytesIO()
    book.save(content)
    content.seek(0)
    monkeypatch.setattr(dashboard.rx, "get_upload_dir", lambda: tmp_path)
    monkeypatch.setattr(mail_watcher, "UPLOADS_FOLDER", tmp_path)
    monkeypatch.setattr(mail_watcher, "load_mail_items", lambda: [])
    monkeypatch.setattr(dashboard, "was_processed", lambda name: None)
    state = upload_state()
    asyncio.run(dashboard.State.handle_upload.fn(state, [UploadFile(content, filename="заказ.xlsx")]))
    saved = tmp_path / "заказ.xlsx"
    assert saved.exists()
    assert mail_watcher.resolve_order_for_download(saved.name) == saved.resolve()
    loaded = load_workbook(saved)
    assert loaded.active.cell(2, 2).value == 12
    loaded.close()
    assert state.selected_file == "заказ.xlsx"
    assert state.uploaded_file_path == str(saved)
    assert state.status == "Файл загружен: заказ.xlsx"
    assert state.order_mail_metadata is None
    assert state.output_file == state.log_file == state.error_text == ""


def test_order_details_offer_processed_and_source_as_distinct_downloads(tmp_path, monkeypatch):
    source_folder = tmp_path / "uploads"
    source_folder.mkdir()
    source = source_folder / "заказ.xlsx"
    original = Workbook()
    original.active.title = "Исходный"
    original.active.append(["Дыня", 5])
    original.save(source)
    original.close()

    result = tmp_path / "processed" / "20.10.26" / "заказ_обработан_12-00-00.xlsx"
    result.parent.mkdir(parents=True)
    processed = Workbook()
    processed.active.title = "Маршрут №4"
    processed.active.append(["Товар", "Сумма", "ФМ 47"])
    processed.active.append(["Дыня", "", 5])
    processed.save(result)
    processed.close()

    monkeypatch.setattr(paths, "PROCESSED_FOLDER", result.parent.parent)
    monkeypatch.setattr(paths, "ORDER_RECORDS_FILE", tmp_path / "order_records.json")
    monkeypatch.setattr(mail_watcher, "UPLOADS_FOLDER", source_folder)
    monkeypatch.setattr(mail_watcher, "load_mail_items", lambda: [])
    order_records.create_processed_order_record(result, source.name, "Клиент", "Город")
    state = SimpleNamespace(_refresh_order_invoices=lambda: None)
    state._show_order_details = dashboard.State._show_order_details.__get__(state)
    dashboard.State.open_order_details.fn(state, source.name, "20.10.26 12:00")

    assert state.selected_order_processed_path == f"20.10.26/{result.name}"
    assert state.selected_order_processed_name == result.name
    downloads = []
    monkeypatch.setattr(dashboard.rx, "download", lambda **kwargs: downloads.append(kwargs))
    dashboard.State.download_order_processed_file.fn(state)
    dashboard.State.download_order_source_file.fn(state)

    client = TestClient(dashboard.custom_api)
    token = dashboard_sessions.create_session()
    try:
        client.cookies.set(dashboard_sessions.SESSION_COOKIE, token)
        result_response = client.get(downloads[0]["url"])
        source_response = client.get(downloads[1]["url"])
    finally:
        dashboard_sessions.destroy_session(token)
    assert result_response.status_code == source_response.status_code == 200
    assert result_response.content == result.read_bytes() != source_response.content
    assert source_response.content == source.read_bytes()
    assert downloads[0]["filename"] == result.name
    assert downloads[1]["filename"] == source.name
    processed_copy = load_workbook(BytesIO(result_response.content))
    source_copy = load_workbook(BytesIO(source_response.content))
    assert processed_copy.sheetnames == ["Маршрут №4"]
    assert source_copy.sheetnames == ["Исходный"]
    processed_copy.close()
    source_copy.close()

    dashboard.State._show_order_details(state, source.name, "20.10.26 11:00", "Получен")
    assert state.selected_order_processed_path == ""
    assert state.selected_order_processed_name == ""


def test_order_details_labels_downloads_unambiguously():
    rendered = str(dashboard.order_details_drawer())
    assert json.dumps("Скачать обработанный") in rendered
    assert json.dumps("Скачать исходный") in rendered
    assert "download_order_processed_file" in rendered
    assert "download_order_source_file" in rendered


def test_missing_processed_result_is_reported_without_serving_source(monkeypatch):
    state = SimpleNamespace(selected_order_processed_path="20.10.26/missing.xlsx", order_detail_status="")
    monkeypatch.setattr(dashboard, "resolve_processed_file", lambda relative: (_ for _ in ()).throw(
        dashboard.ProcessedFileError("Файл не найден")
    ))
    downloads = []
    monkeypatch.setattr(dashboard.rx, "download", lambda **kwargs: downloads.append(kwargs))

    dashboard.State.download_order_processed_file.fn(state)

    assert "Обработанный файл больше недоступен" in state.order_detail_status
    assert downloads == []


def test_manual_order_source_download_requires_login(tmp_path, monkeypatch):
    monkeypatch.setattr(mail_watcher, "UPLOADS_FOLDER", tmp_path)
    monkeypatch.setattr(mail_watcher, "load_mail_items", lambda: [])
    source = tmp_path / "ручной заказ.xlsx"
    source.write_bytes(b"original uploaded workbook")
    client = TestClient(Starlette(routes=[
        Route("/orders/{filename}", download_order_source_file, methods=["GET"]),
    ]))
    assert client.get("/orders/ручной заказ.xlsx").status_code == 401
    token = dashboard_sessions.create_session()
    try:
        client.cookies.set(dashboard_sessions.SESSION_COOKIE, token)
        response = client.get("/orders/ручной заказ.xlsx")
        assert response.status_code == 200
        assert response.content == source.read_bytes()
        assert "attachment" in response.headers["content-disposition"]
    finally:
        dashboard_sessions.destroy_session(token)


@pytest.mark.parametrize("filename", ["../outside.xlsx", "nested/order.xlsx", "", "secret.json", "missing.xlsx"])
def test_manual_order_download_rejects_unavailable_files(tmp_path, monkeypatch, filename):
    monkeypatch.setattr(mail_watcher, "UPLOADS_FOLDER", tmp_path)
    monkeypatch.setattr(mail_watcher, "load_mail_items", lambda: [])
    (tmp_path / "secret.json").write_text("{}", encoding="utf-8")
    with pytest.raises(mail_watcher.MailDownloadError):
        mail_watcher.resolve_order_for_download(filename)


def test_manual_upload_without_selection_has_clear_status():
    state = upload_state()
    asyncio.run(dashboard.State.handle_upload.fn(state, []))
    assert state.status == "Сначала выберите Excel-файл"
    assert state.uploaded_file_path == ""


def test_manual_upload_rejects_non_excel_without_writing(tmp_path, monkeypatch):
    monkeypatch.setattr(dashboard.rx, "get_upload_dir", lambda: tmp_path)
    state = upload_state()
    asyncio.run(dashboard.State.handle_upload.fn(state, [UploadFile(BytesIO(b"text"), filename="notes.txt")]))
    assert "Неподдерживаемый формат" in state.status
    assert not list(tmp_path.iterdir())


def test_upload_button_uses_same_file_context_as_picker():
    def walk(component):
        yield component
        for child in component.children:
            yield from walk(child)

    button = next(component for component in walk(dashboard.order_panel())
                  if component.tag == "Button" and "handle_upload" in str(component.event_triggers))
    rendered = str(button.render())
    assert 'filesById?.["order_upload"]' in rendered
    assert '"uploadFiles"' in rendered
    picker = dashboard.upload_area()
    assert str(picker.children[0].id) == dashboard.UPLOAD_ID
    assert len(picker.children[0].children) == 2  # Native input and one content container.
    stylesheet = (Path(__file__).parents[1] / "assets/dashboard-design.css").read_text(encoding="utf-8")
    assert '.file-picker input[type="file"] { display:block !important;' in stylesheet
