from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from modules import driver_orders, paths


@pytest.fixture
def order_fixture(tmp_path, monkeypatch):
    output = tmp_path / "processed.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Маршрут №1"
    sheet.append(["Товар", "Сумма", "фм 10", "фм 14"])
    sheet.append(["Яблоки", "=SUM(C2:D2)", 10, 4])
    sheet.append(["Бананы", "=SUM(C3:D3)", 2.5, 0])
    sheet.append(["ИТОГО", "=SUM(B2:B3)", None, None])
    workbook.save(output)
    monkeypatch.setattr(paths, "DRIVER_ORDERS_FOLDER", tmp_path / "driver_orders")
    return output


def assignment():
    return {
        "route": "route_1",
        "label": "Маршрут №1",
        "driver_id": "driver-1",
        "driver_name": "Иван",
        "vehicle_id": "vehicle-1",
        "vehicle_name": "Газель",
        "vehicle_plate": "А123ВС71",
    }


def measurement(boxes=1, tare="0", gross="0"):
    return {
        "box_count": str(boxes),
        "tare_weight_per_box_kg": str(tare),
        "gross_weight_kg": str(gross),
    }


def store_measurements(store, gross_by_name):
    return {
        line["line_id"]: measurement(0, 0, gross_by_name[line["name"]])
        for line in store["lines"]
    }


def record_mapping():
    return {
        "id": "record-1",
        "routes": [{
            "route_key": "route_1",
            "destinations": [
                {"id": "destination-10", "name": "фм 10", "positions": [
                    {"id": "position-apples", "name": "Яблоки", "loading_line_id": "loading-apples"},
                    {"id": "position-bananas", "name": "Бананы", "loading_line_id": "loading-bananas"},
                ]},
                {"id": "destination-14", "name": "фм 14", "positions": [
                    {"id": "position-apples-14", "name": "Яблоки", "loading_line_id": "loading-apples"},
                    {"id": "position-bananas-14", "name": "Бананы", "loading_line_id": "loading-bananas"},
                ]},
            ],
        }],
    }


def test_build_snapshot_reads_route_sheet_in_kg_and_links_record(order_fixture):
    snapshot = driver_orders.build_assignment_snapshot(
        order_fixture, "Город", assignment(), order_record=record_mapping()
    )

    assert snapshot["route_key"] == "route_1"
    assert snapshot["order_record_id"] == "record-1"
    assert [store["name"] for store in snapshot["stores"]] == ["фм 10", "фм 14"]
    assert [line["required_qty"] for line in snapshot["stores"][0]["lines"]] == ["10", "2.5"]
    assert [line["unit"] for line in snapshot["stores"][0]["lines"]] == ["kg", "kg"]
    assert snapshot["stores"][0]["order_record_destination_id"] == "destination-10"
    assert snapshot["stores"][0]["lines"][0]["order_record_position_id"] == "position-apples"


def test_completion_stores_measurements_and_is_idempotent(order_fixture):
    snapshot = driver_orders.build_assignment_snapshot(order_fixture, "Город", assignment())
    path = driver_orders.save_snapshot(snapshot)
    store = snapshot["stores"][0]
    values = store_measurements(store, {"Яблоки": "8", "Бананы": "2"})

    result = driver_orders.complete_store(path, store["store_id"], values, "operation-1", 0)

    assert result["revision"] == 1
    assert result["completed_stores"] == 1
    saved_lines = result["stores"][0]["lines"]
    assert [line["actual_qty"] for line in saved_lines] == ["8", "2"]
    assert saved_lines[0]["shipment_measurement"]["net_weight_kg"] == "8"
    summary = {item["name"]: item for item in result["summary"]}
    assert summary["Яблоки"]["completed"] == 8

    repeated = driver_orders.complete_store(path, store["store_id"], values, "operation-1", 0)
    assert repeated["revision"] == 1
    assert repeated["summary"] == result["summary"]


def test_completion_rejects_stale_revision_missing_lines_and_invalid_measurement(order_fixture):
    snapshot = driver_orders.build_assignment_snapshot(order_fixture, "Город", assignment())
    path = driver_orders.save_snapshot(snapshot)
    store = snapshot["stores"][0]
    values = store_measurements(store, {"Яблоки": "1", "Бананы": "1"})
    driver_orders.complete_store(path, store["store_id"], values, "operation-1", 0)

    with pytest.raises(driver_orders.DriverOrderConflict):
        driver_orders.complete_store(path, store["store_id"], values, "operation-2", 0)
    with pytest.raises(driver_orders.DriverOrderError, match="каждой позиции"):
        driver_orders.complete_store(path, snapshot["stores"][1]["store_id"], {}, "operation-3", 1)

    second = snapshot["stores"][1]
    invalid = store_measurements(second, {"Яблоки": "1", "Бананы": "1"})
    invalid[second["lines"][0]["line_id"]] = measurement(2, 1, 1)
    with pytest.raises(driver_orders.DriverOrderError, match="меньше веса тары"):
        driver_orders.complete_store(path, second["store_id"], invalid, "operation-4", 1)


def test_legacy_offline_quantity_is_preserved_as_net_only(order_fixture):
    snapshot = driver_orders.build_assignment_snapshot(order_fixture, "Город", assignment())
    path = driver_orders.save_snapshot(snapshot)
    store = snapshot["stores"][0]
    quantities = {line["line_id"]: "3" for line in store["lines"]}

    result = driver_orders.complete_store(
        path, store["store_id"], quantities, "operation-1", 0, legacy_quantities=True
    )

    saved = result["stores"][0]["lines"][0]["shipment_measurement"]
    assert saved["net_weight_kg"] == "3"
    assert saved["legacy_net_weight_kg"] == "3"
    assert saved["box_count"] == ""


def test_completed_delivery_report_uses_calculated_net_weight(order_fixture):
    snapshot = driver_orders.build_assignment_snapshot(order_fixture, "Город", assignment())
    path = driver_orders.save_snapshot(snapshot)
    store = snapshot["stores"][0]
    values = store_measurements(store, {"Яблоки": "8.25", "Бананы": "2.25"})
    driver_orders.complete_store(path, store["store_id"], values, "operation-1", 0)
    broken = path.parent / ("f" * 32 + ".json")
    broken.write_text("{broken", encoding="utf-8")

    report = driver_orders.completed_delivery_report()
    summary = {item["name"]: item for item in report["summary"]}

    assert report["completed_stores"] == 1
    assert report["line_count"] == 2
    assert summary["Яблоки"] == {
        "name": "Яблоки", "unit": "kg", "planned": 10, "delivered": 8.25, "difference": -1.75
    }


def test_republishing_assignments_preserves_completed_store_without_double_counting(order_fixture):
    first_paths = driver_orders.publish_assignments(order_fixture, "Город", "order.xlsx", [assignment()])
    snapshot = driver_orders.load_snapshot(first_paths[0])
    store = snapshot["stores"][0]
    values = store_measurements(store, {"Яблоки": "1", "Бананы": "1"})
    driver_orders.complete_store(first_paths[0], store["store_id"], values, "operation-1", 0)

    republished_paths = driver_orders.publish_assignments(order_fixture, "Город", "order.xlsx", [assignment()])
    report = driver_orders.completed_delivery_report()

    assert republished_paths == first_paths
    assert driver_orders.assignments_for_vehicle("vehicle-1")[0]["assignment_id"] == snapshot["assignment_id"]
    assert report["completed_stores"] == 1
    assert report["line_count"] == 2


def test_invalid_planned_quantity_reports_processed_sheet_cell_and_value(order_fixture):
    workbook = load_workbook(order_fixture)
    workbook.active["C2"] = "—"
    workbook.save(order_fixture)
    workbook.close()

    with pytest.raises(driver_orders.DriverOrderError) as exc:
        driver_orders.build_assignment_snapshot(order_fixture, "Город", assignment())

    message = str(exc.value)
    assert "Маршрут №1" in message
    assert "C2" in message
    assert "фм 10" in message
    assert "Яблоки" in message
    assert "'—'" in message
    assert "плановое количество" in message


def test_invalid_planned_quantity_cannot_partially_publish_assignments(order_fixture):
    workbook = load_workbook(order_fixture)
    workbook.active["D3"] = "не число"
    workbook.save(order_fixture)
    workbook.close()

    with pytest.raises(driver_orders.DriverOrderError, match="D3"):
        driver_orders.publish_assignments(order_fixture, "Город", "order.xlsx", [assignment()])

    assert driver_orders._snapshots(active_only=False) == []
    assert not driver_orders.index_path().exists()


def test_valid_decimal_comma_still_publishes(order_fixture):
    workbook = load_workbook(order_fixture)
    workbook.active["C2"] = "1,5"
    workbook.save(order_fixture)
    workbook.close()

    snapshot = driver_orders.build_assignment_snapshot(order_fixture, "Город", assignment())
    assert snapshot["stores"][0]["lines"][0]["required_qty"] == "1.5"
