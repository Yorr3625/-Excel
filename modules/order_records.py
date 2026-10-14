"""Durable order snapshots for warehouse loading and store shipment."""

from __future__ import annotations

from copy import deepcopy
from datetime import datetime
from decimal import Decimal, InvalidOperation, localcontext
import json
import os
from pathlib import Path
import tempfile
import threading
import uuid

from openpyxl import load_workbook

from modules import paths
from modules.route_drivers import ROUTE_AREAS


SCHEMA_VERSION = 1
MAX_AMOUNT_DIGITS = 18
MAX_OPERATION_ID_LENGTH = 128
ZERO_REASON_OPTIONS = (
    "Недостаток товара на складе",
    "Брак или несоответствие качеству",
    "Ошибка в заявке",
    "Отсутствие тары",
)
_LOCK = threading.RLock()


class OrderRecordError(ValueError):
    """Order record data is invalid or unavailable."""


class OrderRecordNotFound(OrderRecordError):
    """The requested order record does not exist."""


class OrderRecordConflict(OrderRecordError):
    """The record was changed by another operation."""


def _now() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def _decimal(value, field: str) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise OrderRecordError(f"Некорректное значение: {field}")
    try:
        number = Decimal(str(value).strip().replace(",", "."))
    except (AttributeError, InvalidOperation) as error:
        raise OrderRecordError(f"Некорректное значение: {field}") from error
    if not number.is_finite() or number < 0:
        raise OrderRecordError(f"{field} не может быть отрицательным")
    if len(number.as_tuple().digits) > MAX_AMOUNT_DIGITS:
        raise OrderRecordError(f"Слишком большое значение: {field}")
    return number


def _quantity(value, field: str) -> str:
    number = _decimal(value, field)
    rendered = format(number, "f")
    return rendered.rstrip("0").rstrip(".") if "." in rendered else rendered


def _integer_quantity(value, field: str) -> str:
    number = _decimal(value, field)
    if number != number.to_integral_value():
        raise OrderRecordError(f"{field} должен быть целым числом")
    return str(int(number))


def calculate_net_weight(gross_weight, tare_weight_per_box, box_count) -> Decimal:
    """Return gross weight minus tare weight for every box."""

    return _decimal(gross_weight, "Грязный вес") - (
        _decimal(tare_weight_per_box, "Вес тары")
        * _decimal(box_count, "Количество ящиков")
    )


def _empty_measurement(box_count: str = "") -> dict:
    return {
        "box_count": str(box_count),
        "tare_weight_per_box_kg": "",
        "gross_weight_kg": "",
        "net_weight_kg": "",
        "zero_reason": "",
    }


def _measurement_complete(measurement: object) -> bool:
    return isinstance(measurement, dict) and all(
        str(measurement.get(key, "")).strip() != ""
        for key in ("box_count", "tare_weight_per_box_kg", "gross_weight_kg", "net_weight_kg")
    )


def is_zero_measurement(measurement: object) -> bool:
    if not isinstance(measurement, dict):
        return False
    fields = ("box_count", "tare_weight_per_box_kg", "gross_weight_kg")
    if any(str(measurement.get(field, "")).strip() == "" for field in fields):
        return False
    try:
        return all(_decimal(measurement[field], field) == 0 for field in fields)
    except OrderRecordError:
        return False


def _shipment_measurement_complete(measurement: object) -> bool:
    if isinstance(measurement, dict) and measurement.get("kind") == "net":
        try:
            _validated_shipment_measurement(measurement)
            return True
        except OrderRecordError:
            return False
    if _measurement_complete(measurement):
        return True
    return isinstance(measurement, dict) and str(measurement.get("legacy_net_weight_kg", "")).strip() != ""


def _validated_measurement(value: object, require_zero_reason: bool = False) -> dict:
    if not isinstance(value, dict):
        raise OrderRecordError("Некорректные данные взвешивания")
    box_count = _integer_quantity(value.get("box_count", ""), "Количество ящиков")
    tare_weight = _quantity(value.get("tare_weight_per_box_kg", ""), "Вес тары")
    gross_weight = _quantity(value.get("gross_weight_kg", ""), "Грязный вес")
    net_weight = calculate_net_weight(gross_weight, tare_weight, box_count)
    if net_weight < 0:
        raise OrderRecordError("Грязный вес не может быть меньше веса тары")
    measurement = {
        "box_count": box_count,
        "tare_weight_per_box_kg": tare_weight,
        "gross_weight_kg": gross_weight,
        "net_weight_kg": _quantity(net_weight, "Чистый вес"),
        "zero_reason": "",
    }
    if is_zero_measurement(measurement):
        zero_reason = str(value.get("zero_reason", "")).strip()
        if require_zero_reason and zero_reason not in ZERO_REASON_OPTIONS:
            raise OrderRecordError("Укажите причину нулевой позиции")
        if zero_reason in ZERO_REASON_OPTIONS:
            measurement["zero_reason"] = zero_reason
    return measurement


def loading_measurement_parts(value: object) -> list[dict]:
    value = value if isinstance(value, dict) else {}
    if "weighings" not in value:
        return [value]
    parts = value["weighings"]
    if not isinstance(parts, list) or not parts or len(parts) > 100:
        raise OrderRecordError("Некорректный список взвешиваний")
    ids = set()
    for part in parts:
        if not isinstance(part, dict) or not isinstance(part.get("id"), str) or not part["id"] or part["id"] in ids or "weighings" in part:
            raise OrderRecordError("Некорректные идентификаторы взвешиваний")
        ids.add(part["id"])
    return parts


def loading_measurement_totals(value: object) -> dict:
    """Preview complete portions without treating an unfinished weighing as zero."""
    parts = loading_measurement_parts(value)
    boxes = gross = tare = net = Decimal(0)
    complete = True
    with localcontext() as arithmetic:
        arithmetic.prec = 60
        for part in parts:
            try:
                validated = _validated_measurement(part)
            except OrderRecordError:
                complete = False
                continue
            count = _decimal(validated["box_count"], "Ящики")
            boxes += count
            gross += _decimal(validated["gross_weight_kg"], "Брутто")
            tare += count * _decimal(validated["tare_weight_per_box_kg"], "Тара")
            net += _decimal(validated["net_weight_kg"], "Нетто")
    return {
        "box_count": _format_number(boxes) if complete else "",
        "gross_weight_kg": _format_number(gross) if complete else "",
        "total_tare_weight_kg": _format_number(tare) if complete else "",
        "net_weight_kg": _format_number(net) if complete else "",
        "partial_net_weight_kg": _format_number(net),
        "complete": complete,
    }


def _refresh_loading_measurement(measurement: dict) -> None:
    if "weighings" not in measurement:
        _refresh_draft_net_weight(measurement)
        return
    parts = loading_measurement_parts(measurement)
    for part in parts:
        _refresh_draft_net_weight(part)
    totals = loading_measurement_totals(measurement)
    for field in ("box_count", "gross_weight_kg", "total_tare_weight_kg", "net_weight_kg"):
        measurement[field] = totals[field]
    measurement["tare_weight_per_box_kg"] = ""
    measurement["zero_reason"] = ""


def _validated_loading_measurement(value: object, require_zero_reason: bool = False) -> dict:
    if not isinstance(value, dict) or "weighings" not in value:
        return _validated_measurement(value, require_zero_reason=require_zero_reason)
    parts = []
    for part in loading_measurement_parts(value):
        validated = _validated_measurement(part, require_zero_reason=require_zero_reason)
        if _decimal(part.get("net_weight_kg"), "Нетто") != _decimal(validated["net_weight_kg"], "Нетто"):
            raise OrderRecordError("Нетто взвешивания не согласовано с весами")
        parts.append({"id": part["id"], **validated})
    result = {"weighings": parts}
    _refresh_loading_measurement(result)
    for field in ("box_count", "gross_weight_kg", "total_tare_weight_kg", "net_weight_kg"):
        if _decimal(value.get(field), field) != _decimal(result[field], field):
            raise OrderRecordError("Итог взвешиваний не согласован с отдельными порциями")
    return result


def _loading_measurement_complete(value: object) -> bool:
    if not isinstance(value, dict) or "weighings" not in value:
        return _measurement_complete(value)
    try:
        _validated_loading_measurement(value, require_zero_reason=True)
        return True
    except OrderRecordError:
        return False


def _validated_shipment_measurement(value: object) -> dict:
    if isinstance(value, dict) and value.get("kind") == "net":
        if set(value) != {"kind", "net_weight_kg"}:
            raise OrderRecordError("Для чистого веса передайте только нетто")
        return {"kind": "net", "net_weight_kg": _quantity(value.get("net_weight_kg", ""), "Чистый вес")}
    if isinstance(value, dict) and "kind" in value:
        raise OrderRecordError("Некорректный тип измерения отгрузки")
    if isinstance(value, dict) and "legacy_net_weight_kg" in value:
        net_weight = _quantity(value.get("legacy_net_weight_kg", ""), "Фактическое количество")
        return {
            "box_count": "",
            "tare_weight_per_box_kg": "",
            "gross_weight_kg": "",
            "net_weight_kg": net_weight,
            "legacy_net_weight_kg": net_weight,
        }
    return _validated_measurement(value)


def _actor(value: object) -> dict:
    value = value if isinstance(value, dict) else {}
    return {
        "kind": str(value.get("kind", "")),
        "id": str(value.get("id", "")),
        "name": str(value.get("name", "")),
    }


def _atomic_write(payload: dict) -> None:
    path = paths.ORDER_RECORDS_FILE
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        descriptor, temporary = tempfile.mkstemp(
            prefix=f".{path.name}.", suffix=".tmp", dir=str(path.parent)
        )
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as stream:
            json.dump(payload, stream, ensure_ascii=False, indent=2)
            stream.write("\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
        temporary = None
    except OSError as error:
        raise OrderRecordError("Не удалось сохранить реестр заказов") from error
    finally:
        if temporary:
            try:
                os.unlink(temporary)
            except OSError:
                pass


def _empty_payload() -> dict:
    return {"schema_version": SCHEMA_VERSION, "records": []}


def _line_key(name: object, unit: object) -> tuple[str, str]:
    return str(name or "").strip().casefold(), str(unit or "kg").strip().casefold()


def _normalize_measurement(line: dict) -> None:
    measurement = line.get("measurement")
    if isinstance(measurement, dict):
        if "weighings" in measurement:
            for part in loading_measurement_parts(measurement):
                _normalize_measurement({"measurement": part})
        for key in ("box_count", "tare_weight_per_box_kg", "gross_weight_kg", "net_weight_kg", "zero_reason"):
            measurement[key] = str(measurement.get(key, ""))
        if measurement["zero_reason"] not in ZERO_REASON_OPTIONS:
            measurement["zero_reason"] = ""
        return

    legacy_box_count = line.get("loaded_box_count", line.pop("loading_amount", ""))
    line["measurement"] = _empty_measurement(legacy_box_count)


def _normalize_shipment(position: dict) -> None:
    shipment = position.get("shipment")
    if not isinstance(shipment, dict):
        position["shipment"] = {
            "status": "pending",
            "operation_id": "",
            "completed_at": "",
            "completed_by": _actor({}),
            "measurement": _empty_measurement(),
        }
        return

    shipment["status"] = "completed" if shipment.get("status") == "completed" else "pending"
    shipment["operation_id"] = str(shipment.get("operation_id", ""))
    shipment["completed_at"] = str(shipment.get("completed_at", ""))
    shipment["completed_by"] = _actor(shipment.get("completed_by"))
    measurement = shipment.get("measurement")
    if not isinstance(measurement, dict):
        measurement = {
            "box_count": shipment.get("box_count", ""),
            "tare_weight_per_box_kg": shipment.get("tare_weight_per_box_kg", ""),
            "gross_weight_kg": shipment.get("gross_weight_kg", ""),
            "net_weight_kg": shipment.get("net_weight_kg", ""),
        }
    if measurement.get("kind") == "net":
        shipment["measurement"] = {"kind": "net", "net_weight_kg": str(measurement.get("net_weight_kg", ""))}
        return
    shipment["measurement"] = {
        key: str(measurement.get(key, ""))
        for key in ("box_count", "tare_weight_per_box_kg", "gross_weight_kg", "net_weight_kg")
    }
    if "legacy_net_weight_kg" in measurement:
        shipment["measurement"]["legacy_net_weight_kg"] = str(measurement["legacy_net_weight_kg"])


def _empty_route_loading() -> dict:
    return {
        "mode": "route",
        "status": "new",
        "last_saved_at": "",
        "last_saved_by": _actor({}),
        "confirmed_at": "",
        "confirmed_by": _actor({}),
    }


def _normalize_route_loading(value: object) -> dict:
    value = value if isinstance(value, dict) else {}
    status = str(value.get("status", "new"))
    return {
        "mode": "stores" if value.get("mode") == "stores" else "route",
        "status": status if status in {"new", "draft", "partial", "loaded"} else "new",
        "last_saved_at": str(value.get("last_saved_at", "")),
        "last_saved_by": _actor(value.get("last_saved_by")),
        "confirmed_at": str(value.get("confirmed_at", "")),
        "confirmed_by": _actor(value.get("confirmed_by")),
        "batches": deepcopy(value.get("batches", [])),
    }


def _route_keys(record: dict) -> list[str]:
    return [
        str(route.get("route_key", ""))
        for route in record.get("routes", [])
        if isinstance(route, dict) and str(route.get("route_key", ""))
    ]


def _route_lines(record: dict, route_key: str) -> list[dict]:
    return [
        line
        for line in record.get("loading_lines", [])
        if isinstance(line, dict) and str(line.get("route_key", "")) == route_key
    ]


def _legacy_loading_ready(record: dict) -> bool:
    return (
        (record.get("loading") or {}).get("status") == "loaded"
        and bool(record.get("loading_lines"))
        and all(_loading_measurement_complete(line.get("measurement")) for line in record["loading_lines"])
    )


def route_loading_mode(record: dict, route_key: str) -> str:
    state = ((record.get("loading") or {}).get("route_states") or {}).get(route_key, {})
    return "stores" if state.get("mode") == "stores" else "route"


def _store_loading_destinations(record: dict, route_key: str) -> list[dict]:
    routes = [route for route in record.get("routes", []) if route.get("route_key") == route_key]
    lines = _route_lines(record, route_key)
    line_ids = {str(line.get("id", "")) for line in lines}
    if len(routes) != 1 or not lines or "" in line_ids or len(line_ids) != len(lines):
        raise OrderRecordError("Нет надёжной связи товаров с маршрутом для раздельной загрузки")
    destinations = routes[0].get("destinations", [])
    lines_by_id = {str(line["id"]): line for line in lines}
    destination_ids, position_ids, used_lines = set(), set(), set()
    if not destinations:
        raise OrderRecordError("В маршруте отсутствуют магазины")
    for destination in destinations:
        destination_id = str(destination.get("id", ""))
        positions = destination.get("positions", [])
        if not destination_id or destination_id in destination_ids or not positions:
            raise OrderRecordError("Некорректные связи магазинов маршрута")
        destination_ids.add(destination_id)
        for position in positions:
            position_id = str(position.get("id", ""))
            line_id = str(position.get("loading_line_id", ""))
            if not position_id or position_id in position_ids or line_id not in line_ids:
                raise OrderRecordError("Нет точной связи позиции магазина с загрузкой маршрута")
            if _line_key(position.get("name"), position.get("unit")) != _line_key(lines_by_id[line_id].get("name"), lines_by_id[line_id].get("unit")):
                raise OrderRecordError("Товар магазина не соответствует связанной позиции маршрута")
            position_ids.add(position_id)
            used_lines.add(line_id)
    if used_lines != line_ids:
        raise OrderRecordError("Не все товары маршрута связаны с магазинами")
    return destinations


def store_loading_available(record: dict, route_key: str) -> bool:
    try:
        _store_loading_destinations(record, route_key)
    except OrderRecordError:
        return False
    return True


def store_loading_measurements(record: dict, route_key: str, line_id: str) -> list[dict]:
    return [
        (position.get("loading") or {}).get("measurement") or {}
        for destination in _store_loading_destinations(record, route_key)
        for position in destination.get("positions", [])
        if position.get("loading_line_id") == line_id
    ]


def loading_batches(record: dict, route_key: str) -> list[dict]:
    batches = ((record.get("loading") or {}).get("route_states") or {}).get(route_key, {}).get("batches", [])
    if not isinstance(batches, list):
        raise OrderRecordError("Некорректные отправки загрузки")
    valid_stores = {str(destination.get("id", "")) for route in record.get("routes", [])
                    if route.get("route_key") == route_key for destination in route.get("destinations", [])}
    seen_ids, seen_stores, seen_operations = set(), set(), set()
    for batch in batches:
        if (not isinstance(batch, dict) or not isinstance(batch.get("id"), str) or not batch["id"]
                or batch["id"] in seen_ids or not isinstance(batch.get("operation_id"), str)
                or not batch["operation_id"] or batch["operation_id"] in seen_operations
                or not isinstance(batch.get("source_assignment_id"), str) or not batch["source_assignment_id"]):
            raise OrderRecordError("Некорректные отправки загрузки")
        stores = batch.get("destination_ids")
        if not isinstance(stores, list) or not stores or any(not isinstance(identity, str) or not identity for identity in stores):
            raise OrderRecordError("Некорректный состав отправки")
        if len(set(stores)) != len(stores) or seen_stores.intersection(stores) or not set(stores).issubset(valid_stores):
            raise OrderRecordError("Магазин включён в несколько отправок или не относится к маршруту")
        seen_ids.add(batch["id"])
        seen_operations.add(batch["operation_id"])
        seen_stores.update(stores)
    return batches


def loading_batch_destinations(record: dict, route_key: str, batch_id: str) -> list[dict]:
    if route_loading_mode(record, route_key) != "stores":
        raise OrderRecordConflict("Частичная отправка доступна только при загрузке по магазинам")
    batch = next((item for item in loading_batches(record, route_key) if item["id"] == batch_id), None)
    if batch is None:
        raise OrderRecordNotFound("Отправка загрузки не найдена")
    destinations = _store_loading_destinations(record, route_key)
    selected = [item for item in destinations if item["id"] in batch["destination_ids"]]
    if len(selected) != len(batch["destination_ids"]) or any(
        (item.get("loading") or {}).get("status") != "loaded"
        or not all(_store_measurement_ready((position.get("loading") or {}).get("measurement")) for position in item["positions"])
        for item in selected
    ):
        raise OrderRecordConflict("Состав или измерения отправки изменились")
    return selected


def release_store_loading(order_record_id: str, revision: int, route_key: str,
                          source_assignment_id: str, operation_id: str, actor: dict) -> dict:
    if not isinstance(operation_id, str) or not operation_id or len(operation_id) > MAX_OPERATION_ID_LENGTH:
        raise OrderRecordError("Некорректный идентификатор операции")
    if not isinstance(source_assignment_id, str) or not source_assignment_id:
        raise OrderRecordError("Нет водительского назначения маршрута")
    if not isinstance(actor, dict) or actor.get("kind") not in {"loader", "dashboard"} or not actor.get("id"):
        raise OrderRecordError("Нет доступа к отправке загрузки")
    with _LOCK:
        payload = _read_payload()
        record = _find_record(payload, order_record_id)
        batches = loading_batches(record, route_key)
        known = next((item for item in batches if item.get("operation_id") == operation_id), None)
        if known:
            if known.get("source_assignment_id") != source_assignment_id:
                raise OrderRecordConflict("Операция относится к другому назначению")
            return deepcopy(known)
        _check_loading_revision(record, revision)
        if route_loading_mode(record, route_key) != "stores":
            raise OrderRecordConflict("Частичная отправка доступна только при загрузке по магазинам")
        destinations = _store_loading_destinations(record, route_key)
        released = {identity for batch in batches for identity in batch["destination_ids"]}
        selected = [item for item in destinations if item["id"] not in released and (item.get("loading") or {}).get("status") == "loaded"]
        if not selected:
            raise OrderRecordConflict("Нет новых подтверждённых магазинов для отправки")
        for destination in selected:
            if any((position.get("shipment") or {}).get("status") == "completed" for position in destination["positions"]):
                raise OrderRecordConflict("Магазин уже отгружен")
            for position in destination["positions"]:
                _validated_loading_measurement((position.get("loading") or {}).get("measurement"), require_zero_reason=True)
        batch = {"id": uuid.uuid4().hex, "source_assignment_id": source_assignment_id,
                 "destination_ids": [item["id"] for item in selected], "operation_id": operation_id,
                 "created_at": _now(), "created_by": _actor(actor)}
        _route_state(record, route_key).setdefault("batches", []).append(batch)
        record["revision"] += 1
        _refresh_overall_loading_status(record)
        _atomic_write(payload)
        return deepcopy(batch)


def loading_line_net_weight(record: dict, line: dict) -> Decimal:
    """Confirmed route totals share the original loading-line identity in both modes."""
    route_key = str(line.get("route_key", ""))
    batches = loading_batches(record, route_key)
    if batches:
        released = {identity for batch in batches for identity in batch["destination_ids"]}
        destinations = [destination for destination in _store_loading_destinations(record, route_key) if destination["id"] in released]
        if len(destinations) != len(released) or any(
            (destination.get("loading") or {}).get("status") != "loaded"
            or not all(_store_measurement_ready((position.get("loading") or {}).get("measurement")) for position in destination["positions"])
            for destination in destinations
        ):
            raise OrderRecordConflict("Состав или измерения отправки изменились")
        measurements = [(position.get("loading") or {}).get("measurement") or {}
                        for destination in destinations for position in destination["positions"]
                        if position.get("loading_line_id") == line["id"]]
    else:
        if not (route_loading_ready(record, route_key) if route_key else _legacy_loading_ready(record)):
            return Decimal(0)
        measurements = store_loading_measurements(record, route_key, str(line["id"])) if route_loading_mode(record, route_key) == "stores" else []
    if route_loading_mode(record, route_key) == "stores":
        with localcontext() as arithmetic:
            arithmetic.prec = 60
            return sum((
                _decimal(_validated_loading_measurement(measurement, require_zero_reason=True)["net_weight_kg"], "Нетто")
                for measurement in measurements
            ), Decimal(0))
    measurement = line.get("measurement") or {}
    return _decimal(measurement["net_weight_kg"], "Нетто") if _loading_measurement_complete(measurement) else Decimal(0)


def _store_measurement_ready(measurement: object) -> bool:
    if not _loading_measurement_complete(measurement):
        return False
    try:
        validated = _validated_loading_measurement(measurement, require_zero_reason=True)
        return _decimal(measurement["net_weight_kg"], "Нетто") == _decimal(validated["net_weight_kg"], "Нетто")
    except OrderRecordError:
        return False


def route_loading_ready(record: dict, route_key: str) -> bool:
    lines = _route_lines(record, route_key)
    if not lines:
        return _legacy_loading_ready(record)
    state = ((record.get("loading") or {}).get("route_states") or {}).get(route_key, {})
    if route_loading_mode(record, route_key) == "stores":
        try:
            destinations = _store_loading_destinations(record, route_key)
        except OrderRecordError:
            return False
        return state.get("status") == "loaded" and all(
            (destination.get("loading") or {}).get("status") == "loaded"
            and all(_store_measurement_ready((position.get("loading") or {}).get("measurement")) for position in destination.get("positions", []))
            for destination in destinations
        )
    return state.get("status") == "loaded" and all(
        _loading_measurement_complete(line.get("measurement")) for line in lines
    )


def route_return_balance(record: dict, route_key: str, loading_batch_id: str = "") -> dict:
    """Keep exact destination facts to detect incomplete two-file shipment saves."""
    routes = [route for route in record.get("routes", []) if route.get("route_key") == route_key]
    lines = _route_lines(record, route_key)
    if len(routes) != 1 or not lines:
        raise OrderRecordError("Нет надёжной связи загрузки с маршрутом")
    if loading_batch_id:
        destinations = loading_batch_destinations(record, route_key, loading_batch_id)
        selected_line_ids = {position["loading_line_id"] for destination in destinations for position in destination["positions"]}
        lines = [line for line in lines if line["id"] in selected_line_ids]
    else:
        if loading_batches(record, route_key):
            raise OrderRecordConflict("Для частичной загрузки выберите конкретную отправку")
        if not route_loading_ready(record, route_key):
            raise OrderRecordError("Склад ещё не подтвердил загрузку маршрута")
        destinations = routes[0].get("destinations", [])

    def number(value, field):
        if len(str(value)) > 64:
            raise OrderRecordError(f"Слишком большое значение: {field}")
        result = _decimal(value, field)
        if abs(result.as_tuple().exponent) > 9 or result > Decimal("1000000000"):
            raise OrderRecordError(f"Слишком большое значение: {field}")
        return result

    def net(measurement, legacy=False):
        if isinstance(measurement, dict) and "weighings" in measurement and not legacy:
            validated = _validated_loading_measurement(measurement, require_zero_reason=True)
            return sum((net(part) for part in validated["weighings"]), Decimal(0))
        if not isinstance(measurement, dict):
            raise OrderRecordError("Отсутствует подтверждённое измерение нетто")
        result = number(measurement.get("net_weight_kg"), "Нетто")
        if legacy and measurement.get("kind") == "net":
            return number(_validated_shipment_measurement(measurement)["net_weight_kg"], "Нетто")
        if legacy and "legacy_net_weight_kg" in measurement:
            expected = number(measurement["legacy_net_weight_kg"], "Нетто")
        else:
            boxes = number(measurement.get("box_count"), "Количество ящиков")
            if boxes != boxes.to_integral_value():
                raise OrderRecordError("Количество ящиков должно быть целым")
            tare = number(measurement.get("tare_weight_per_box_kg"), "Вес тары")
            gross = number(measurement.get("gross_weight_kg"), "Грязный вес")
            expected = calculate_net_weight(gross, tare, boxes)
        if result != expected:
            raise OrderRecordError("Измерение нетто не согласовано с взвешиванием")
        return result

    with localcontext() as arithmetic:
        arithmetic.prec = 60
        balances = {}
        for line in lines:
            line_id = str(line.get("id", ""))
            if not line_id or line_id in balances:
                raise OrderRecordError("Некорректные идентификаторы загрузки")
            positions = [position for destination in destinations for position in destination["positions"] if position.get("loading_line_id") == line_id]
            balances[line_id] = {
                "loading_line_id": line_id, "name": str(line.get("name", "")),
                "planned": sum((number(position.get("planned_amount"), "Заказано") for position in positions), Decimal(0)) if loading_batch_id else number(line.get("planned_amount"), "Заказано"),
                "loaded": sum((net((position.get("loading") or {}).get("measurement")) for position in positions), Decimal(0))
                if route_loading_mode(record, route_key) == "stores" else net(line.get("measurement")),
                "shipped": Decimal("0"),
            }
        destination_facts = []
        destination_ids, position_ids, used_lines = set(), set(), set()
        complete = True
        for destination in destinations:
            destination_id = str(destination.get("id", ""))
            if not destination_id or destination_id in destination_ids or not destination.get("positions"):
                raise OrderRecordError("Некорректные связи магазинов маршрута")
            destination_ids.add(destination_id)
            facts = []
            for position in destination["positions"]:
                position_id = str(position.get("id", ""))
                line_id = str(position.get("loading_line_id", ""))
                if not position_id or position_id in position_ids or line_id not in balances:
                    raise OrderRecordError("Нет надёжной связи отгрузки с загрузкой маршрута")
                position_ids.add(position_id)
                used_lines.add(line_id)
                shipment = position.get("shipment") or {}
                completed = shipment.get("status") == "completed"
                shipped = net(shipment.get("measurement"), legacy=True) if completed else Decimal("0")
                for visit in shipment.get("extra", []):
                    shipped += net(visit.get("measurement"), legacy=True)
                balances[line_id]["shipped"] += shipped
                complete = complete and completed
                facts.append({"id": position_id, "loading_line_id": line_id,
                              "completed": completed, "net_weight_kg": _format_number(shipped) if completed else ""})
            destination_facts.append({"id": destination_id, "positions": sorted(facts, key=lambda item: item["id"])})
        if not destination_facts or used_lines != set(balances):
            raise OrderRecordError("Нет надёжной связи магазинов с загрузкой маршрута")
        result_lines = [{
            "loading_line_id": line["loading_line_id"], "name": line["name"],
            "planned_kg": _format_number(line["planned"]),
            "loaded_net_kg": _format_number(line["loaded"]),
            "shipped_net_kg": _format_number(line["shipped"]),
            "available_net_kg": _format_number(line["loaded"] - line["shipped"]),
        } for line in balances.values()]
    return {"route_key": route_key, "route_label": str(routes[0].get("route_label", "")),
            "complete": complete, "lines": sorted(result_lines, key=lambda item: item["loading_line_id"]),
            "destinations": sorted(destination_facts, key=lambda item: item["id"])}


def _refresh_overall_loading_status(record: dict) -> None:
    route_keys = _route_keys(record)
    if not any(_route_lines(record, route_key) for route_key in route_keys):
        return
    loading = record["loading"]
    states = loading["route_states"]
    for route_key in route_keys:
        if loading_batches(record, route_key) and states[route_key]["status"] != "loaded":
            states[route_key]["status"] = "partial"
    if route_keys and all(route_loading_ready(record, route_key) for route_key in route_keys):
        loading["status"] = "loaded"
    elif any(states[route_key]["status"] == "partial" for route_key in route_keys):
        loading["status"] = "partial"
    elif any(states[route_key]["status"] != "new" for route_key in route_keys):
        loading["status"] = "draft"
    else:
        loading["status"] = "new"


def _normalize_legacy_record(record: dict) -> None:
    loading_lines = record.get("loading_lines")
    if not isinstance(loading_lines, list):
        record["loading_lines"] = []
        loading_lines = record["loading_lines"]

    route_keys = set(_route_keys(record))
    line_ids_by_key: dict[tuple[str, str], str] = {}
    route_line_ids_by_key: dict[tuple[str, str, str], str] = {}
    for line in loading_lines:
        if not isinstance(line, dict):
            continue
        if line.get("unit") in {"", "box", None}:
            line["unit"] = "kg"
        line.setdefault("id", uuid.uuid4().hex)
        line.setdefault("planned_amount", "0")
        route_key = str(line.get("route_key", ""))
        if route_key not in route_keys:
            line.pop("route_key", None)
            route_key = ""
        _normalize_measurement(line)
        key = _line_key(line.get("name"), line.get("unit"))
        line_ids_by_key[key] = str(line["id"])
        if route_key:
            route_line_ids_by_key[(route_key, *key)] = str(line["id"])

    for route in record.get("routes", []):
        if not isinstance(route, dict):
            continue
        route_key = str(route.get("route_key", ""))
        for destination in route.get("destinations", []):
            if not isinstance(destination, dict):
                continue
            destination.setdefault("id", uuid.uuid4().hex)
            for position in destination.get("positions", []):
                if not isinstance(position, dict):
                    continue
                if position.get("unit") in {"", "box", None}:
                    position["unit"] = "kg"
                position.setdefault("id", uuid.uuid4().hex)
                position.setdefault("planned_amount", "0")
                line_key = _line_key(position.get("name"), position.get("unit"))
                position.setdefault(
                    "loading_line_id",
                    route_line_ids_by_key.get(
                        (route_key, *line_key), line_ids_by_key.get(line_key, "")
                    ),
                )
                _normalize_shipment(position)

    loading = record.get("loading")
    loading = loading if isinstance(loading, dict) else {}
    states = loading.get("route_states")
    states = states if isinstance(states, dict) else {}
    record["loading"] = {
        "status": str(loading.get("status", "new")),
        "last_saved_at": str(loading.get("last_saved_at", "")),
        "last_saved_by": _actor(loading.get("last_saved_by")),
        "route_states": {
            route_key: _normalize_route_loading(states.get(route_key))
            for route_key in route_keys
        },
    }
    if record["loading"]["status"] not in {"new", "draft", "partial", "loaded"}:
        record["loading"]["status"] = "new"
    _refresh_overall_loading_status(record)
    record.setdefault("revision", 0)


def _read_payload() -> dict:
    path = paths.ORDER_RECORDS_FILE
    if not path.exists():
        return _empty_payload()
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        raise OrderRecordError("Реестр заказов повреждён или недоступен") from error
    if (
        not isinstance(payload, dict)
        or payload.get("schema_version") != SCHEMA_VERSION
        or not isinstance(payload.get("records"), list)
    ):
        raise OrderRecordError("Неподдерживаемая версия реестра заказов")
    for record in payload["records"]:
        if not isinstance(record, dict):
            raise OrderRecordError("Реестр заказов повреждён")
        _normalize_legacy_record(record)
    return payload


def _validated_relative_output_path(output_file: str | Path) -> str:
    root = Path(paths.PROCESSED_FOLDER).resolve(strict=False)
    output = Path(output_file).resolve(strict=False)
    try:
        relative = output.relative_to(root)
    except ValueError as error:
        raise OrderRecordError("Результат заказа находится вне каталога обработки") from error
    if len(relative.parts) != 2 or output.suffix.casefold() != ".xlsx" or not output.is_file():
        raise OrderRecordError("Результат заказа недоступен")
    return relative.as_posix()


def _route_sheet_number(title: str) -> int | None:
    prefix = "Маршрут №"
    if not title.startswith(prefix):
        return None
    try:
        number = int(title.removeprefix(prefix))
    except ValueError:
        return None
    return number if 1 <= number <= 99 else None


def _store_columns(sheet) -> list[tuple[int, str]]:
    columns = []
    for column in range(3, sheet.max_column + 1):
        name = str(sheet.cell(1, column).value or "").strip()
        if name and name.casefold() != "сумма":
            columns.append((column, name))
    return columns


def _aggregate_line(aggregate: dict[tuple[str, str], dict], name: str) -> dict:
    key = _line_key(name, "kg")
    return aggregate.setdefault(
        key,
        {
            "id": uuid.uuid4().hex,
            "name": name,
            "unit": "kg",
            "planned_amount": Decimal("0"),
        },
    )


def _route_snapshot(sheet, route_number: int) -> tuple[dict, list[dict]]:
    route_key = f"route_{route_number}"
    aggregate: dict[tuple[str, str], dict] = {}
    stores = _store_columns(sheet)
    destinations = [
        {"id": uuid.uuid4().hex, "name": name, "positions": []}
        for _, name in stores
    ]
    destination_by_column = {
        column: destination for (column, _), destination in zip(stores, destinations)
    }

    for row in range(2, sheet.max_row + 1):
        name = str(sheet.cell(row, 1).value or "").strip()
        if not name or name.casefold() == "итого":
            continue
        for column, _ in stores:
            value = sheet.cell(row, column).value
            if value in (None, ""):
                continue
            planned_amount = _quantity(value, "Заказано")
            aggregate_line = _aggregate_line(aggregate, name)
            aggregate_line["planned_amount"] += _decimal(planned_amount, "Заказано")
            destination_by_column[column]["positions"].append(
                {
                    "id": uuid.uuid4().hex,
                    "loading_line_id": aggregate_line["id"],
                    "name": name,
                    "unit": "kg",
                    "planned_amount": planned_amount,
                    "shipment": {
                        "status": "pending",
                        "operation_id": "",
                        "completed_at": "",
                        "completed_by": _actor({}),
                        "measurement": _empty_measurement(),
                    },
                }
            )

    route = {
        "route_key": route_key,
        "route_label": f"Маршрут №{route_number}",
        "area": ROUTE_AREAS[route_number - 1] if route_number <= len(ROUTE_AREAS) else "",
        "destinations": [destination for destination in destinations if destination["positions"]],
    }
    lines = [
        {
            "id": line["id"],
            "route_key": route_key,
            "name": line["name"],
            "unit": line["unit"],
            "planned_amount": _quantity(line["planned_amount"], "Заказано"),
            "measurement": _empty_measurement(),
        }
        for line in aggregate.values()
    ]
    lines.sort(key=lambda item: (item["name"].casefold(), item["unit"]))
    return route, lines


def _workbook_snapshot(output_file: Path) -> tuple[list[dict], list[dict]]:
    try:
        workbook = load_workbook(output_file, data_only=True, read_only=True)
    except (OSError, ValueError) as error:
        raise OrderRecordError("Не удалось прочитать обработанный заказ") from error

    routes = []
    lines = []
    try:
        sheets = sorted(
            (
                (number, sheet)
                for sheet in workbook.worksheets
                if (number := _route_sheet_number(sheet.title)) is not None
            ),
            key=lambda item: item[0],
        )
        for number, sheet in sheets:
            route, route_lines = _route_snapshot(sheet, number)
            if route["destinations"]:
                routes.append(route)
                lines.extend(route_lines)
    finally:
        workbook.close()

    if not routes or not lines:
        raise OrderRecordError("В обработанном заказе не найдены позиции маршрутов")
    return routes, lines


def _mail_metadata(metadata: dict | None) -> dict | None:
    if metadata is None:
        return None
    return {
        key: str(metadata.get(key, ""))
        for key in (
            "message_id",
            "sender",
            "subject",
            "received_at",
            "received_display",
            "source_name",
            "source_email",
            "source_person",
        )
    }


def create_processed_order_record(
    output_file: str | Path,
    source_filename: str,
    client_name: str,
    mode: str,
    mail_metadata: dict | None = None,
) -> dict:
    source_filename = Path(str(source_filename)).name
    client_name = str(client_name).strip()
    if not source_filename or not client_name:
        raise OrderRecordError("Укажите номер заказа и клиента")
    if mode not in {"Город", "Область"}:
        raise OrderRecordError("Не указан режим заказа")

    relative_output_path = _validated_relative_output_path(output_file)
    routes, loading_lines = _workbook_snapshot(Path(output_file))
    record = {
        "id": uuid.uuid4().hex,
        "schema_version": SCHEMA_VERSION,
        "revision": 0,
        "source_filename": source_filename,
        "order_number": Path(source_filename).stem,
        "client_name": client_name,
        "mail_metadata": _mail_metadata(mail_metadata),
        "processed_at": _now(),
        "output_relative_path": relative_output_path,
        "mode": mode,
        "loading": {
            "status": "new",
            "last_saved_at": "",
            "last_saved_by": _actor({}),
            "route_states": {
                str(route["route_key"]): _empty_route_loading()
                for route in routes
            },
        },
        "routes": routes,
        "loading_lines": loading_lines,
    }
    with _LOCK:
        payload = _read_payload()
        payload["records"].append(record)
        _atomic_write(payload)
    return deepcopy(record)


def _record_loading_ready(record: dict) -> bool:
    route_keys = _route_keys(record)
    if any(_route_lines(record, route_key) for route_key in route_keys):
        return bool(route_keys) and all(
            route_loading_ready(record, route_key) for route_key in route_keys
        )
    return _legacy_loading_ready(record)


def loading_route_summaries(record: dict) -> list[dict]:
    route_labels = {
        str(route.get("route_key", "")): str(route.get("route_label", ""))
        for route in record.get("routes", [])
        if isinstance(route, dict)
    }
    summaries = []
    for route_key in _route_keys(record):
        lines = _route_lines(record, route_key)
        if not lines:
            continue
        state = ((record.get("loading") or {}).get("route_states") or {}).get(route_key, {})
        summaries.append(
            {
                "route_key": route_key,
                "route_label": route_labels.get(route_key, route_key),
                "status": str(state.get("status", "new")),
                "line_count": len(lines),
                "ready": route_loading_ready(record, route_key),
            }
        )
    return summaries


def _destinations(record: dict):
    for route in record.get("routes", []):
        if not isinstance(route, dict):
            continue
        for destination in route.get("destinations", []):
            if isinstance(destination, dict):
                yield route, destination


def shipment_status(record: dict) -> str:
    if not _record_loading_ready(record):
        if any((position.get("shipment") or {}).get("status") == "completed" for _, destination in _destinations(record) for position in destination.get("positions", [])):
            return "shipment_in_progress"
        if any(loading_batches(record, route_key) for route_key in _route_keys(record)):
            return "ready_for_shipment"
        return "waiting_loading"
    destinations = list(_destinations(record))
    if not destinations:
        return "shipment_completed"
    completed = [
        all((position.get("shipment") or {}).get("status") == "completed" for position in destination.get("positions", []))
        for _, destination in destinations
    ]
    if all(completed):
        return "shipment_completed"
    return "shipment_in_progress" if any(completed) else "ready_for_shipment"


def _format_number(value: Decimal) -> str:
    rendered = format(value, "f")
    return rendered.rstrip("0").rstrip(".") if "." in rendered else rendered


def logistics_summary(record: dict) -> dict:
    lines: dict[tuple[str, str], dict] = {}
    line_groups: dict[str, tuple[str, str]] = {}
    for loading_line in record.get("loading_lines", []):
        if not isinstance(loading_line, dict):
            continue
        key = _line_key(loading_line.get("name"), loading_line.get("unit"))
        line = lines.setdefault(
            key,
            {
                "id": str(loading_line.get("id", "")),
                "name": str(loading_line.get("name", "")),
                "planned": Decimal("0"),
                "loaded": Decimal("0"),
                "shipped": Decimal("0"),
            },
        )
        line_groups[str(loading_line.get("id", ""))] = key
        try:
            line["planned"] += _decimal(loading_line.get("planned_amount", "0"), "Заказано")
        except OrderRecordError:
            continue
        route_key = str(loading_line.get("route_key", ""))
        ready = route_loading_ready(record, route_key) if route_key else _legacy_loading_ready(record)
        if ready or loading_batches(record, route_key):
            line["loaded"] += loading_line_net_weight(record, loading_line)

    for _, destination in _destinations(record):
        for position in destination.get("positions", []):
            if not isinstance(position, dict):
                continue
            shipment = position.get("shipment") or {}
            measurement = shipment.get("measurement") or {}
            line = lines.get(line_groups.get(str(position.get("loading_line_id", ""))))
            if line is None or shipment.get("status") != "completed" or not _shipment_measurement_complete(measurement):
                continue
            line["shipped"] += _decimal(measurement["net_weight_kg"], "Чистый вес")
            for visit in shipment.get("extra", []):
                added = _validated_shipment_measurement(visit["measurement"])
                line["shipped"] += _decimal(added["net_weight_kg"], "Чистый вес")

    items = []
    total_ordered = Decimal("0")
    total_loaded = Decimal("0")
    total_shipped = Decimal("0")
    total_shortage = Decimal("0")
    total_surplus = Decimal("0")
    for line in sorted(lines.values(), key=lambda item: item["name"].casefold()):
        shortage = max(line["planned"] - line["shipped"], Decimal("0"))
        surplus = max(line["shipped"] - line["planned"], Decimal("0"))
        items.append(
            {
                "id": line["id"],
                "name": line["name"],
                "planned_kg": _format_number(line["planned"]),
                "loaded_net_kg": _format_number(line["loaded"]),
                "shipped_net_kg": _format_number(line["shipped"]),
                "underload_kg": _format_number(max(line["planned"] - line["loaded"], Decimal("0"))),
                "shortage_kg": _format_number(shortage),
                "surplus_kg": _format_number(surplus),
                "loading_to_shipment_kg": _format_number(line["loaded"] - line["shipped"]),
            }
        )
        total_ordered += line["planned"]
        total_loaded += line["loaded"]
        total_shipped += line["shipped"]
        total_shortage += shortage
        total_surplus += surplus

    return {
        "loading_ready": _record_loading_ready(record),
        "shipment_status": shipment_status(record),
        "routes": loading_route_summaries(record),
        "items": items,
        "total_ordered_kg": _format_number(total_ordered),
        "total_loaded_net_kg": _format_number(total_loaded),
        "total_shipped_net_kg": _format_number(total_shipped),
        "total_underload_kg": _format_number(max(total_ordered - total_loaded, Decimal("0"))),
        "total_shortage_kg": _format_number(total_shortage),
        "total_surplus_kg": _format_number(total_surplus),
        "total_loading_to_shipment_kg": _format_number(total_loaded - total_shipped),
    }


def list_route_trip_summaries(processed_date: str, mode: str) -> list[dict]:
    requested_date = datetime.strptime(processed_date, "%Y-%m-%d").date()
    with _LOCK:
        records = _read_payload()["records"]
    trips = []
    for record in records:
        if record.get("mode") != mode:
            continue
        try:
            processed_at = datetime.fromisoformat(str(record.get("processed_at", "")))
        except ValueError:
            continue
        if processed_at.astimezone().date() != requested_date:
            continue
        for route in record.get("routes", []):
            route_key = str(route.get("route_key", ""))
            lines = _route_lines(record, route_key)
            linked = bool(lines)
            batches = loading_batches(record, route_key)
            released = {identity for batch in batches for identity in batch["destination_ids"]}
            ready = linked and route_loading_ready(record, route_key) and all(
                item["id"] in released for item in route.get("destinations", [])
            ) if batches else linked and route_loading_ready(record, route_key)
            loading_status = ((record.get("loading") or {}).get("route_states") or {}).get(route_key, {}).get("status", "new")
            stores = []
            shipped_total = Decimal(0)
            confirmed_positions = 0
            for destination in route.get("destinations", []):
                positions = destination.get("positions", [])
                confirmed = [
                    position for position in positions
                    if (position.get("shipment") or {}).get("status") == "completed"
                    and _shipment_measurement_complete((position.get("shipment") or {}).get("measurement"))
                ]
                confirmed_positions += len(confirmed)
                shipped = sum((
                    _decimal(position["shipment"]["measurement"]["net_weight_kg"], "Отгружено")
                    for position in confirmed
                ), Decimal(0))
                shipped_total += shipped
                completed = bool(positions) and len(confirmed) == len(positions)
                stores.append({
                    "id": str(destination.get("id", "")), "name": str(destination.get("name", "")),
                    "completed": completed,
                    "status": "Отгружено" if completed else "Частично" if confirmed else "Ожидает",
                    "shipped_net_kg": _format_number(shipped) if confirmed else "—",
                })
            completed_count = sum(store["completed"] for store in stores)
            if stores and completed_count == len(stores):
                stage = "Отгрузка завершена"
            elif confirmed_positions:
                stage = "Отгрузка идёт"
            elif ready:
                stage = "Готов к отгрузке"
            elif batches or loading_status == "partial":
                stage = "Загружен частично"
            elif loading_status == "draft":
                stage = "Загрузка начата"
            else:
                stage = "Ожидает загрузки"
            loaded = sum((loading_line_net_weight(record, line) for line in lines), Decimal(0)) if ready or loading_batches(record, route_key) else None
            trips.append({
                "order_record_id": str(record["id"]), "order_number": str(record.get("order_number", "")),
                "client_name": str(record.get("client_name", "")), "processed_at": processed_at.astimezone().isoformat(),
                "mode": mode, "route_key": route_key, "route_label": str(route.get("route_label", route_key)),
                "stage": stage, "loading_ready": ready, "limited_data": not linked,
                "loaded_net_kg": _format_number(loaded) if loaded is not None else "—",
                "shipped_net_kg": _format_number(shipped_total) if confirmed_positions else "—",
                "completed_stores": completed_count, "store_count": len(stores), "stores": stores,
            })
    return sorted(trips, key=lambda trip: (trip["processed_at"], trip["order_record_id"], trip["route_key"]), reverse=True)


def list_order_record_summaries() -> list[dict]:
    with _LOCK:
        records = _read_payload()["records"]
    summaries = []
    for record in records:
        summary = logistics_summary(record)
        summaries.append(
            {
                "id": str(record.get("id", "")),
                "order_number": str(record.get("order_number", "")),
                "client_name": str(record.get("client_name", "")),
                "processed_at": str(record.get("processed_at", "")),
                "status": str((record.get("loading") or {}).get("status", "new")),
                "shipment_status": summary["shipment_status"],
            }
        )
    return sorted(summaries, key=lambda item: item["processed_at"], reverse=True)


def latest_record_ids_by_source_filename() -> dict[str, str]:
    with _LOCK:
        records = _read_payload()["records"]
    latest = {}
    for record in sorted(records, key=lambda item: str(item.get("processed_at", ""))):
        filename = str(record.get("source_filename", ""))
        record_id = str(record.get("id", ""))
        if filename and record_id:
            latest[filename] = record_id
    return latest


def load_order_record(order_record_id: str) -> dict:
    with _LOCK:
        for record in _read_payload()["records"]:
            if record.get("id") == order_record_id:
                return deepcopy(record)
    raise OrderRecordNotFound("Заказ не найден")


def _find_record(payload: dict, order_record_id: str) -> dict:
    record = next(
        (item for item in payload["records"] if item.get("id") == order_record_id),
        None,
    )
    if record is None:
        raise OrderRecordNotFound("Заказ не найден")
    return record


def _draft_value(value: object, field: str) -> str:
    if value in (None, ""):
        return ""
    if field == "box_count":
        return _integer_quantity(value, "Количество ящиков")
    if field == "tare_weight_per_box_kg":
        return _quantity(value, "Вес тары")
    if field == "gross_weight_kg":
        return _quantity(value, "Грязный вес")
    zero_reason = str(value).strip()
    if zero_reason not in ZERO_REASON_OPTIONS:
        raise OrderRecordError("Некорректная причина нулевой позиции")
    return zero_reason


def _refresh_draft_net_weight(measurement: dict) -> None:
    measurement["net_weight_kg"] = ""
    if any(
        str(measurement.get(field, "")).strip() == ""
        for field in ("box_count", "tare_weight_per_box_kg", "gross_weight_kg")
    ):
        return
    try:
        net_weight = calculate_net_weight(
            measurement["gross_weight_kg"],
            measurement["tare_weight_per_box_kg"],
            measurement["box_count"],
        )
    except OrderRecordError:
        return
    if net_weight >= 0:
        measurement["net_weight_kg"] = _quantity(net_weight, "Чистый вес")


def _route_state(record: dict, route_key: str) -> dict:
    states = record["loading"]["route_states"]
    if route_key not in states:
        raise OrderRecordNotFound("Маршрут не найден в заказе")
    return states[route_key]


def _check_loading_revision(record: dict, revision: int) -> None:
    if type(revision) is not int or revision < 0:
        raise OrderRecordError("Некорректная версия заказа")
    if record.get("revision") != revision:
        raise OrderRecordConflict("Данные загрузки уже изменены другим пользователем")


def loading_mode_locked(record: dict, route_key: str) -> bool:
    state = ((record.get("loading") or {}).get("route_states") or {}).get(route_key, {})
    if state.get("status", "new") != "new":
        return True
    if any(any(str(value) for value in (line.get("measurement") or {}).values()) for line in _route_lines(record, route_key)):
        return True
    for route, destination in _destinations(record):
        if route.get("route_key") != route_key:
            continue
        if (destination.get("loading") or {}).get("status", "new") != "new":
            return True
        for position in destination.get("positions", []):
            if (position.get("shipment") or {}).get("status", "pending") == "completed":
                return True
            if any(str(value) for value in ((position.get("loading") or {}).get("measurement") or {}).values()):
                return True
    return False


def set_route_loading_mode(order_record_id: str, revision: int, route_key: str, mode: str) -> dict:
    if mode not in {"route", "stores"}:
        raise OrderRecordError("Некорректный режим загрузки")
    with _LOCK:
        payload = _read_payload()
        record = _find_record(payload, order_record_id)
        _check_loading_revision(record, revision)
        state = _route_state(record, route_key)
        if route_loading_mode(record, route_key) == mode:
            return deepcopy(record)
        if loading_mode_locked(record, route_key):
            raise OrderRecordConflict("Режим нельзя изменить после начала ввода или подтверждения загрузки")
        if mode == "stores":
            _store_loading_destinations(record, route_key)
        state["mode"] = mode
        record["revision"] += 1
        _atomic_write(payload)
        return deepcopy(record)


def _loading_part(measurement: dict, weighing_id: str) -> dict:
    if "weighings" not in measurement:
        if weighing_id:
            raise OrderRecordConflict("Взвешивание больше не существует")
        return measurement
    if not weighing_id:
        raise OrderRecordConflict("Обновите форму: у товара несколько взвешиваний")
    part = next((item for item in loading_measurement_parts(measurement) if item["id"] == weighing_id), None)
    if part is None:
        raise OrderRecordConflict("Взвешивание больше не существует")
    return part


def _edit_loading_measurement(source: dict, field: str, value: object, weighing_id: str) -> dict:
    measurement = deepcopy(source) if source else _empty_measurement()
    part = _loading_part(measurement, weighing_id)
    part[field] = _draft_value(value, field)
    if field == "zero_reason" and not is_zero_measurement(part):
        raise OrderRecordError("Причину можно указать только для нулевой позиции")
    if not is_zero_measurement(part):
        part["zero_reason"] = ""
    _refresh_loading_measurement(measurement)
    return measurement


def change_loading_weighing(order_record_id: str, revision: int, route_key: str,
                           destination_id: str, position_id: str, actor: dict,
                           remove_id: str | None = None) -> dict:
    """Atomically add or remove a portion in the exact loading context."""
    with _LOCK:
        payload = _read_payload()
        record = _find_record(payload, order_record_id)
        _check_loading_revision(record, revision)
        state = _route_state(record, route_key)
        if state["status"] == "loaded":
            raise OrderRecordConflict("Загрузка маршрута уже подтверждена")
        store_state = None
        if destination_id:
            if route_loading_mode(record, route_key) != "stores":
                raise OrderRecordConflict("Не выбран режим загрузки по магазинам")
            _store_loading_destinations(record, route_key)
            _, destination = _find_destination(record, route_key, destination_id)
            store_state = destination.setdefault("loading", _empty_route_loading())
            if store_state.get("status") == "loaded":
                raise OrderRecordConflict("Загрузка магазина уже подтверждена")
            position = next((item for item in destination["positions"] if item["id"] == position_id), None)
            if position is None:
                raise OrderRecordNotFound("Позиция не найдена в магазине")
            target = position.setdefault("loading", {})
        else:
            if route_loading_mode(record, route_key) == "stores":
                raise OrderRecordConflict("Выберите магазин")
            target = next((item for item in _route_lines(record, route_key) if item["id"] == position_id), None)
            if target is None:
                raise OrderRecordNotFound("Позиция не найдена в маршруте")
        measurement = deepcopy(target.get("measurement") or _empty_measurement())
        parts = loading_measurement_parts(measurement)
        if remove_id is not None:
            if "weighings" not in measurement or len(parts) == 1:
                raise OrderRecordError("Последнее взвешивание нельзя удалить")
            _loading_part(measurement, remove_id)
            measurement["weighings"] = [item for item in parts if item["id"] != remove_id]
        else:
            if len(parts) >= 100:
                raise OrderRecordError("Максимум 100 взвешиваний для товара")
            if "weighings" not in measurement:
                measurement = {"weighings": [{**measurement, "id": uuid.uuid4().hex}]}
            measurement["weighings"].append({**_empty_measurement(), "id": uuid.uuid4().hex})
        _refresh_loading_measurement(measurement)
        target["measurement"] = measurement
        stamp = {"status": "draft", "last_saved_at": _now(), "last_saved_by": _actor(actor)}
        state.update(stamp)
        if store_state is not None:
            store_state.update(stamp)
        record["revision"] += 1
        _refresh_overall_loading_status(record)
        _atomic_write(payload)
        return deepcopy(record)


def save_store_loading_draft(
    order_record_id: str, route_key: str, destination_id: str,
    position_id: str, field: str, value: object, actor: dict, weighing_id: str = "",
) -> dict:
    if field not in {"box_count", "tare_weight_per_box_kg", "gross_weight_kg", "zero_reason"}:
        raise OrderRecordError("Некорректное поле загрузки")
    with _LOCK:
        payload = _read_payload()
        record = _find_record(payload, order_record_id)
        state = _route_state(record, route_key)
        if route_loading_mode(record, route_key) != "stores":
            raise OrderRecordConflict("Для маршрута не выбран режим загрузки по магазинам")
        _store_loading_destinations(record, route_key)
        _, destination = _find_destination(record, route_key, destination_id)
        store_state = destination.setdefault("loading", _empty_route_loading())
        if state["status"] == "loaded" or store_state.get("status") == "loaded":
            raise OrderRecordConflict("Загрузка магазина уже подтверждена")
        position = next((item for item in destination["positions"] if item.get("id") == position_id), None)
        if position is None:
            raise OrderRecordNotFound("Позиция не найдена в магазине")
        source = (position.get("loading") or {}).get("measurement") or {}
        measurement = _edit_loading_measurement(source, field, value, weighing_id)
        position["loading"] = {"measurement": measurement}
        stamp = {"status": "draft", "last_saved_at": _now(), "last_saved_by": _actor(actor)}
        store_state.update(stamp)
        state.update(stamp)
        record["revision"] += 1
        _refresh_overall_loading_status(record)
        _atomic_write(payload)
        return deepcopy(record)


def confirm_store_loading(
    order_record_id: str, revision: int, route_key: str, destination_id: str, actor: dict,
) -> dict:
    with _LOCK:
        payload = _read_payload()
        record = _find_record(payload, order_record_id)
        _check_loading_revision(record, revision)
        state = _route_state(record, route_key)
        if route_loading_mode(record, route_key) != "stores":
            raise OrderRecordConflict("Для маршрута не выбран режим загрузки по магазинам")
        destinations = _store_loading_destinations(record, route_key)
        _, destination = _find_destination(record, route_key, destination_id)
        store_state = destination.setdefault("loading", _empty_route_loading())
        if state["status"] == "loaded" or store_state.get("status") == "loaded":
            raise OrderRecordConflict("Загрузка магазина уже подтверждена")
        for position in destination["positions"]:
            measurement = _validated_loading_measurement((position.get("loading") or {}).get("measurement"), require_zero_reason=True)
            position["loading"] = {"measurement": measurement}
        stamp = {"last_saved_at": _now(), "last_saved_by": _actor(actor)}
        store_state.update({**stamp, "status": "loaded", "confirmed_at": stamp["last_saved_at"], "confirmed_by": _actor(actor)})
        if all((item.get("loading") or {}).get("status") == "loaded" for item in destinations):
            state.update({**stamp, "status": "loaded", "confirmed_at": stamp["last_saved_at"], "confirmed_by": _actor(actor)})
        else:
            state.update({**stamp, "status": "draft"})
        record["revision"] += 1
        _refresh_overall_loading_status(record)
        _atomic_write(payload)
        return deepcopy(record)


def save_loading_draft(
    order_record_id: str,
    route_key: str,
    line_id: str,
    field: str,
    value: object,
    actor: dict,
    weighing_id: str = "",
) -> dict:
    if field not in {"box_count", "tare_weight_per_box_kg", "gross_weight_kg", "zero_reason"}:
        raise OrderRecordError("Некорректное поле загрузки")

    with _LOCK:
        payload = _read_payload()
        record = _find_record(payload, order_record_id)
        route_state = _route_state(record, route_key)
        if route_loading_mode(record, route_key) == "stores":
            raise OrderRecordConflict("В этом режиме загрузка подтверждается отдельно по магазинам")
        if route_state["status"] == "loaded":
            raise OrderRecordConflict("Загрузка маршрута уже подтверждена")
        line = next(
            (item for item in _route_lines(record, route_key) if item.get("id") == line_id),
            None,
        )
        if line is None:
            raise OrderRecordNotFound("Позиция не найдена в маршруте")

        source_measurement = line.get("measurement") if isinstance(line.get("measurement"), dict) else {}
        measurement = _edit_loading_measurement(source_measurement, field, value, weighing_id)
        line["measurement"] = measurement
        route_state.update(
            {
                "status": "draft",
                "last_saved_at": _now(),
                "last_saved_by": _actor(actor),
            }
        )
        record["revision"] += 1
        _refresh_overall_loading_status(record)
        _atomic_write(payload)
        return deepcopy(record)


def confirm_route_loading(
    order_record_id: str,
    revision: int,
    route_key: str,
    actor: dict,
) -> dict:
    if not isinstance(revision, int) or isinstance(revision, bool) or revision < 0:
        raise OrderRecordError("Некорректная версия заказа")

    with _LOCK:
        payload = _read_payload()
        record = _find_record(payload, order_record_id)
        if record.get("revision") != revision:
            raise OrderRecordConflict("Данные загрузки уже изменены другим пользователем")
        route_state = _route_state(record, route_key)
        if route_loading_mode(record, route_key) == "stores":
            raise OrderRecordConflict("В этом режиме загрузка подтверждается отдельно по магазинам")
        if route_state["status"] == "loaded":
            raise OrderRecordConflict("Загрузка маршрута уже подтверждена")
        lines = _route_lines(record, route_key)
        if not lines:
            raise OrderRecordNotFound("В маршруте отсутствуют позиции загрузки")

        for line in lines:
            line["measurement"] = _validated_loading_measurement(
                line.get("measurement"), require_zero_reason=True
            )

        confirmed_at = _now()
        route_state.update(
            {
                "status": "loaded",
                "last_saved_at": confirmed_at,
                "last_saved_by": _actor(actor),
                "confirmed_at": confirmed_at,
                "confirmed_by": _actor(actor),
            }
        )
        record["revision"] += 1
        _refresh_overall_loading_status(record)
        _atomic_write(payload)
        return deepcopy(record)


def _find_destination(record: dict, route_key: str, destination_id: str) -> tuple[dict, dict]:
    for route, destination in _destinations(record):
        if route.get("route_key") == route_key and destination.get("id") == destination_id:
            return route, destination
    raise OrderRecordNotFound("Магазин не найден в заказе")


def save_store_shipment(
    order_record_id: str,
    route_key: str,
    destination_id: str,
    measurements: dict[str, object],
    operation_id: str,
    actor: dict,
    loading_batch_id: str = "",
) -> dict:
    if not isinstance(operation_id, str) or not operation_id or len(operation_id) > MAX_OPERATION_ID_LENGTH:
        raise OrderRecordError("Некорректный идентификатор операции")
    if not isinstance(measurements, dict):
        raise OrderRecordError("Некорректные данные отгрузки")

    with _LOCK:
        payload = _read_payload()
        record = _find_record(payload, order_record_id)
        if loading_batch_id:
            allowed = {item["id"] for item in loading_batch_destinations(record, route_key, loading_batch_id)}
            if destination_id not in allowed:
                raise OrderRecordConflict("Магазин не включён в эту отправку загрузки")
        elif loading_batches(record, route_key):
            raise OrderRecordConflict("Для частичной загрузки выберите конкретную отправку")
        elif not route_loading_ready(record, route_key):
            raise OrderRecordConflict("Склад ещё не подтвердил загрузку маршрута")
        _, destination = _find_destination(record, route_key, destination_id)
        positions = destination.get("positions")
        if not isinstance(positions, list):
            raise OrderRecordError("В магазине отсутствуют позиции")
        position_ids = {str(position.get("id", "")) for position in positions}
        if not position_ids or set(measurements) != position_ids:
            raise OrderRecordError("Заполните отгрузку для каждой позиции")

        existing_operations = {
            str((position.get("shipment") or {}).get("operation_id", ""))
            for position in positions
        }
        if operation_id in existing_operations:
            return deepcopy(record)
        if any((position.get("shipment") or {}).get("status") == "completed" for position in positions):
            raise OrderRecordConflict("Магазин уже подтверждён")

        completed_at = _now()
        completed_by = _actor(actor)
        for position in positions:
            position["shipment"] = {
                "status": "completed",
                "operation_id": operation_id,
                "completed_at": completed_at,
                "completed_by": completed_by,
                "measurement": _validated_shipment_measurement(measurements[str(position["id"])]),
            }

        record["revision"] += 1
        _atomic_write(payload)
        return deepcopy(record)


def append_store_shipment(order_record_id: str, route_key: str, destination_id: str,
                          measurements: dict[str, object], operation_id: str, actor: dict) -> dict:
    """Record a later visit without overwriting the first delivery."""
    if not isinstance(operation_id, str) or not operation_id or len(operation_id) > MAX_OPERATION_ID_LENGTH:
        raise OrderRecordError("Некорректный идентификатор операции")
    with _LOCK:
        payload = _read_payload()
        record = _find_record(payload, order_record_id)
        _, destination = _find_destination(record, route_key, destination_id)
        positions = destination.get("positions", [])
        if not isinstance(measurements, dict) or set(measurements) != {str(position["id"]) for position in positions}:
            raise OrderRecordError("Заполните отгрузку для каждой позиции")
        validated = {identity: _validated_shipment_measurement(value) for identity, value in measurements.items()}
        if any((position.get("shipment") or {}).get("status") != "completed" for position in positions):
            raise OrderRecordConflict("Первую поставку магазина нужно записать отдельно")
        if any(position["shipment"].get("operation_id") == operation_id for position in positions):
            raise OrderRecordConflict("Операция уже записана как первая поставка")
        present = [next((item for item in (position.get("shipment") or {}).get("extra", []) if item.get("operation_id") == operation_id), None)
                   for position in positions]
        if any(present):
            if not all(present) or any(item["measurement"] != validated[str(position["id"])] for position, item in zip(positions, present)):
                raise OrderRecordConflict("Повторная поставка содержит другие данные")
            return deepcopy(record)
        when = _now()
        by = _actor(actor)
        for position in positions:
            position["shipment"].setdefault("extra", []).append({"operation_id": operation_id, "completed_at": when,
                                                                 "completed_by": by, "measurement": validated[str(position["id"])]})
        record["revision"] += 1
        _atomic_write(payload)
        return deepcopy(record)
