"""Sidecar orders used by the authenticated driver workspace."""

from __future__ import annotations

import hashlib
import json
import os
import tempfile
import threading
import uuid
from copy import deepcopy
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from modules import paths


SCHEMA_VERSION = 1
INDEX_SCHEMA_VERSION = 1
MAX_OPERATION_ID_LENGTH = 128
MAX_QUANTITY_DIGITS = 18


class DriverOrderError(ValueError):
    """Invalid driver order data or operation."""


class DriverOrderNotFound(DriverOrderError):
    """The requested sidecar or entity does not exist."""


class DriverOrderConflict(DriverOrderError):
    """The client is working from an obsolete revision."""


_LOCK = threading.RLock()


def _now() -> str:
    return datetime.now().isoformat(timespec="seconds")


def _decimal(value: Any, field: str = "quantity") -> Decimal:
    if isinstance(value, bool) or value is None:
        raise DriverOrderError(f"Некорректное значение {field}")
    try:
        number = Decimal(str(value).strip().replace(",", "."))
    except (InvalidOperation, AttributeError) as error:
        raise DriverOrderError(f"Некорректное значение {field}") from error
    if not number.is_finite() or number < 0:
        raise DriverOrderError(f"{field} не может быть отрицательным")
    if len(number.as_tuple().digits) > MAX_QUANTITY_DIGITS:
        raise DriverOrderError(f"Слишком большое значение {field}")
    return number


def _quantity(value: Any) -> str:
    number = _decimal(value)
    text = format(number, "f")
    if "." in text:
        text = text.rstrip("0").rstrip(".")
    return text or "0"


def _optional_quantity(value: Any) -> str:
    if value in (None, ""):
        return ""
    return _quantity(value)


def _measurement(value: Any) -> dict:
    if not isinstance(value, dict):
        raise DriverOrderError("Некорректные данные отгрузки")
    if value.get("kind") == "net":
        if set(value) != {"kind", "net_weight_kg"}:
            raise DriverOrderError("Для чистого веса передайте только нетто")
        return {"kind": "net", "net_weight_kg": _quantity(value.get("net_weight_kg", ""))}
    if "kind" in value:
        raise DriverOrderError("Некорректный тип измерения отгрузки")
    try:
        box_count = _quantity(value.get("box_count", ""))
        boxes = _decimal(box_count, "количество ящиков")
        if boxes != boxes.to_integral_value():
            raise DriverOrderError("Количество ящиков должно быть целым числом")
        tare = _quantity(value.get("tare_weight_per_box_kg", ""))
        gross = _quantity(value.get("gross_weight_kg", ""))
    except DriverOrderError:
        raise
    net = _decimal(gross, "грязный вес") - boxes * _decimal(tare, "вес тары")
    if net < 0:
        raise DriverOrderError("Грязный вес не может быть меньше веса тары")
    return {
        "box_count": str(int(boxes)),
        "tare_weight_per_box_kg": tare,
        "gross_weight_kg": gross,
        "net_weight_kg": _quantity(net),
    }


def _legacy_measurement(value: Any) -> dict:
    net_weight = _quantity(value)
    return {
        "box_count": "",
        "tare_weight_per_box_kg": "",
        "gross_weight_kg": "",
        "net_weight_kg": net_weight,
        "legacy_net_weight_kg": net_weight,
    }


def _sum(values: list[str]) -> Decimal:
    return sum((_decimal(value) for value in values), Decimal("0"))


def _public_number(value: Decimal) -> int | float:
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def _atomic_write(path: Path, payload: dict) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        descriptor, temporary = tempfile.mkstemp(
            prefix=f".{path.name}.", suffix=".tmp", dir=str(path.parent)
        )
        with os.fdopen(descriptor, "w", encoding="utf-8", newline="\n") as stream:
            json.dump(payload, stream, ensure_ascii=False, indent=2)
            stream.write("\n")
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temporary, path)
        temporary = None
    except OSError as error:
        raise DriverOrderError("Не удалось сохранить заказ водителя") from error
    finally:
        if temporary:
            try:
                os.unlink(temporary)
            except OSError:
                pass


def _read(path: Path) -> dict:
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError) as error:
        raise DriverOrderError("Файл заказа водителя повреждён или недоступен") from error
    if not isinstance(payload, dict) or payload.get("schema_version") != SCHEMA_VERSION:
        raise DriverOrderError("Неподдерживаемая версия заказа водителя")
    return payload


def _empty_index() -> dict:
    return {"schema_version": INDEX_SCHEMA_VERSION, "active": {}, "publications": []}


def index_path() -> Path:
    return paths.DRIVER_ORDERS_FOLDER / "index.json"


def _read_index() -> dict | None:
    path = index_path()
    if not path.exists():
        return None
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, UnicodeError, json.JSONDecodeError):
        return _empty_index()
    if (
        not isinstance(payload, dict)
        or payload.get("schema_version") != INDEX_SCHEMA_VERSION
        or not isinstance(payload.get("active"), dict)
        or not isinstance(payload.get("publications"), list)
    ):
        return _empty_index()
    return payload


def _active_assignment_ids() -> set[str] | None:
    index = _read_index()
    if index is None:
        return None
    return {
        str(assignment_id)
        for assignment_id in index["active"].values()
        if isinstance(assignment_id, str)
    }


def _snapshots(active_only: bool = False) -> list[dict]:
    folder = paths.DRIVER_ORDERS_FOLDER
    if not folder.exists():
        return []
    active_ids = _active_assignment_ids() if active_only else None
    candidates = active_ids if active_ids is not None else [
        path.stem for path in folder.glob("*.json") if path.name != "index.json"
    ]
    snapshots = []
    for assignment_id in candidates:
        try:
            snapshots.append(_read(sidecar_path(str(assignment_id))))
        except DriverOrderError:
            continue
    return snapshots


def _file_digest(path: Path) -> str:
    digest = hashlib.sha256()
    try:
        with path.open("rb") as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b""):
                digest.update(chunk)
    except OSError as error:
        raise DriverOrderError("Результат заказа недоступен") from error
    return digest.hexdigest()


def _stable_id(*parts: str) -> str:
    return hashlib.sha256("\x1f".join(parts).encode("utf-8")).hexdigest()[:24]


def _route_number(route_key: str) -> int:
    try:
        number = int(route_key.removeprefix("route_"))
    except (AttributeError, ValueError) as error:
        raise DriverOrderError("Некорректный маршрут") from error
    if not 1 <= number <= 99:
        raise DriverOrderError("Некорректный маршрут")
    return number


def _sheet_store_columns(sheet) -> list[tuple[int, str]]:
    columns = []
    for column in range(3, sheet.max_column + 1):
        name = str(sheet.cell(1, column).value or "").strip()
        if name and name.casefold() != "сумма":
            columns.append((column, name))
    return columns


def _is_blank_planned_quantity(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _source_rows(sheet, store_columns: list[tuple[int, str]]) -> list[tuple[int, str]]:
    rows = []
    for row in range(2, sheet.max_row + 1):
        name = str(sheet.cell(row, 1).value or "").strip()
        if not name or name.casefold() == "итого":
            continue
        has_numeric = False
        for column, store in store_columns:
            cell = sheet.cell(row, column)
            value = cell.value
            if not _is_blank_planned_quantity(value):
                try:
                    _decimal(value, "плановое количество")
                except DriverOrderError as error:
                    shown_value = repr(value)
                    if len(shown_value) > 80:
                        shown_value = shown_value[:77] + "..."
                    raise DriverOrderError(
                        f"Обработанный лист «{sheet.title}», ячейка {cell.coordinate} "
                        f"(магазин «{store}», товар «{name}»): "
                        f"{error} — {shown_value}"
                    ) from error
                has_numeric = True
        if has_numeric:
            rows.append((row, name))
    return rows


def build_assignment_snapshot(
    output_file: str | Path,
    mode: str,
    assignment: dict,
    order_id: str | None = None,
    order_record: dict | None = None,
) -> dict:
    """Build one driver's route snapshot from a processed route workbook."""
    output_path = Path(output_file)
    route_key = str(assignment.get("route", ""))
    route_number = _route_number(route_key)
    if not output_path.exists():
        raise DriverOrderError("Результат заказа не найден")

    order_id = order_id or _file_digest(output_path)
    assignment_id = _stable_id(
        order_id,
        route_key,
        str(assignment.get("driver_id", "")),
        str(assignment.get("vehicle_id", "")),
    )
    try:
        workbook = load_workbook(output_path, data_only=True, read_only=True)
        sheet = workbook[f"Маршрут №{route_number}"]
    except (OSError, KeyError, ValueError) as error:
        raise DriverOrderError("В результате нет листа назначенного маршрута") from error

    store_columns = _sheet_store_columns(sheet)
    rows = _source_rows(sheet, store_columns)
    stores = []
    for column, store_name in store_columns:
        store_id = _stable_id(order_id, route_key, "store", store_name)
        lines = []
        for row, product_name in rows:
            value = sheet.cell(row, column).value
            if _is_blank_planned_quantity(value):
                continue
            lines.append({
                "line_id": _stable_id(order_id, route_key, store_name, str(row), product_name),
                "name": product_name,
                "unit": "kg",
                "required_qty": _quantity(value),
                "actual_qty": "",
            })
        if not lines:
            continue
        stores.append({
            "store_id": store_id,
            "name": store_name,
            "status": "pending",
            "completed_at": "",
            "photo": "",
            "lines": lines,
        })

    record_routes = {
        str(route.get("route_key", "")): route
        for route in (order_record or {}).get("routes", [])
        if isinstance(route, dict)
    }
    record_route = record_routes.get(route_key, {})
    record_destinations = {
        str(destination.get("name", "")).casefold(): destination
        for destination in record_route.get("destinations", [])
        if isinstance(destination, dict)
    }
    for store in stores:
        destination = record_destinations.get(str(store["name"]).casefold(), {})
        positions = {
            str(position.get("name", "")).casefold(): position
            for position in destination.get("positions", [])
            if isinstance(position, dict)
        }
        store["order_record_destination_id"] = str(destination.get("id", ""))
        for line in store["lines"]:
            position = positions.get(str(line["name"]).casefold(), {})
            line["order_record_position_id"] = str(position.get("id", ""))
            line["loading_line_id"] = str(position.get("loading_line_id", ""))

    return {
        "schema_version": SCHEMA_VERSION,
        "order_id": order_id,
        "order_record_id": str((order_record or {}).get("id", "")),
        "assignment_id": assignment_id,
        "source_file": str(assignment.get("order_file", "")),
        "output_file": str(output_path),
        "mode": str(mode or ""),
        "route_key": route_key,
        "route_label": str(assignment.get("label", f"Маршрут №{route_number}")),
        "driver_id": str(assignment.get("driver_id", "")),
        "driver_name": str(assignment.get("driver_name", "")),
        "vehicle_id": str(assignment.get("vehicle_id", "")),
        "vehicle_name": str(assignment.get("vehicle_name", "")),
        "vehicle_plate": str(assignment.get("vehicle_plate", "")),
        "created_at": _now(),
        "revision": 0,
        "operations": [],
        "stores": stores,
    }


def sidecar_path(assignment_id: str) -> Path:
    if not assignment_id or any(character not in "0123456789abcdef" for character in assignment_id):
        raise DriverOrderError("Некорректный идентификатор назначения")
    return paths.DRIVER_ORDERS_FOLDER / f"{assignment_id}.json"


def save_snapshot(snapshot: dict) -> Path:
    assignment_id = str(snapshot.get("assignment_id", ""))
    path = sidecar_path(assignment_id)
    with _LOCK:
        _atomic_write(path, snapshot)
    return path


def publish_assignments(
    output_file: str | Path,
    mode: str,
    source_file: str,
    assignments: list[dict],
    order_record: dict | None = None,
) -> list[Path]:
    order_id = _file_digest(Path(output_file))
    publication_id = uuid.uuid4().hex
    created_at = _now()
    prepared = []
    for assignment in assignments:
        enriched = {**assignment, "order_file": source_file}
        snapshot = build_assignment_snapshot(output_file, mode, enriched, order_id, order_record)
        logical_assignment_id = _stable_id(str(source_file), str(mode or ""), snapshot["route_key"])
        assignment_id = _stable_id(logical_assignment_id, publication_id)
        snapshot.update({
            "assignment_id": assignment_id,
            "logical_assignment_id": logical_assignment_id,
            "publication_id": publication_id,
            "created_at": created_at,
        })
        prepared.append(snapshot)

    with _LOCK:
        index = _read_index() or _empty_index()
        active = dict(index["active"])
        from modules import vehicle_trips
        for snapshot in prepared:
            existing_id = active.get(snapshot["logical_assignment_id"])
            trip = vehicle_trips.trip_for_assignment(existing_id) if existing_id else None
            if trip and trip["status"] != "closed":
                raise DriverOrderConflict("Маршрут участвует в рейсе машины. Завершите рейс перед повторной публикацией")
        paths_written = []
        publication_assignment_ids = []
        for snapshot in prepared:
            logical_assignment_id = snapshot["logical_assignment_id"]
            existing_id = active.get(logical_assignment_id)
            try:
                existing = _read(sidecar_path(str(existing_id))) if existing_id else None
            except DriverOrderError:
                existing = None
            if existing and (existing.get("loading_batch_id") or any(store.get("status") == "completed" for store in existing.get("stores", []))):
                paths_written.append(sidecar_path(str(existing_id)))
                publication_assignment_ids.append(str(existing_id))
                continue

            path = sidecar_path(snapshot["assignment_id"])
            _atomic_write(path, snapshot)
            active[logical_assignment_id] = snapshot["assignment_id"]
            paths_written.append(path)
            publication_assignment_ids.append(snapshot["assignment_id"])

        publications = list(index["publications"])
        publications.append({
            "publication_id": publication_id,
            "created_at": created_at,
            "assignment_ids": publication_assignment_ids,
            "logical_assignment_ids": [snapshot["logical_assignment_id"] for snapshot in prepared],
        })
        _atomic_write(index_path(), {
            "schema_version": INDEX_SCHEMA_VERSION,
            "active": active,
            "publications": publications,
        })
    return paths_written


def load_snapshot(path_or_id: str | Path) -> dict:
    path = sidecar_path(str(path_or_id)) if not isinstance(path_or_id, Path) else path_or_id
    with _LOCK:
        return _read(path)


def _loading_assignment(order_record_id: str, route_key: str) -> dict:
    candidates = [item for item in _snapshots(active_only=True)
                  if item.get("order_record_id") == order_record_id and item.get("route_key") == route_key]
    current = max(candidates, key=lambda item: item.get("created_at", ""), default=None)
    if current is None or not current.get("driver_id"):
        raise DriverOrderConflict("Маршруту ещё не назначен водитель")
    return current


def _partial_loading_blocker(current: dict) -> str:
    from modules import goods_returns, vehicle_trips
    trip = vehicle_trips.trip_for_assignment(current["assignment_id"])
    if trip:
        if trip.get("projection_error"):
            return "Сначала синхронизируйте отгрузки предыдущей отправки"
        if trip["status"] in {"active", "on_base", "released_pending", "closed"}:
            return ""
        return "Сначала начните или завершите рейс машины предыдущей отправки"
    if current.get("loading_batch_id"):
        if not current.get("stores") or any(store.get("status") != "completed" for store in current["stores"]):
            return "Сначала завершите отгрузку магазинов предыдущей отправки"
        if not trip:
            context = goods_returns.return_context(current)
            if context.get("error") and not context.get("complete"):
                return context["error"]
            document = context.get("document")
            if document and document["status"] != "accepted":
                return "Сначала завершите проверку возврата предыдущей отправки"
            accepted = {line["loading_line_id"]: Decimal(line["measurement"]["net_weight_kg"])
                        for line in (document or {}).get("lines", [])} if document else {}
            if any(Decimal(line["available_net_kg"]) != accepted.get(line["loading_line_id"], Decimal(0)) for line in context["lines"]):
                return "Сначала оформите и завершите сверку остатка предыдущей отправки"
    elif any(store.get("status") == "completed" for store in current.get("stores", [])):
        return "Частичную загрузку нельзя выделить из уже начатой полной отгрузки"
    return ""


def loading_batch_published(batch: dict) -> bool:
    identity = _stable_id(batch["source_assignment_id"], "loading-batch", batch["id"])
    index = _read_index()
    if not index or not any(identity in publication["assignment_ids"] for publication in index["publications"]):
        return False
    try:
        snapshot = _read(sidecar_path(identity))
    except DriverOrderError:
        return False
    return (snapshot.get("assignment_id") == identity and snapshot.get("loading_batch_id") == batch["id"]
            and isinstance(snapshot.get("stores"), list) and len(snapshot["stores"]) == len(batch["destination_ids"])
            and all(isinstance(store, dict) for store in snapshot["stores"])
            and {store.get("order_record_destination_id") for store in snapshot["stores"]} == set(batch["destination_ids"]))


def partial_loading_blocker(order_record_id: str, route_key: str) -> str:
    try:
        return _partial_loading_blocker(_loading_assignment(order_record_id, route_key))
    except (DriverOrderError, ValueError) as error:
        return str(error)


def release_partial_loading(order_record_id: str, route_key: str, revision: int,
                            actor: dict, operation_id: str) -> dict:
    """Publish one immutable store group; retries repair a split record/index write."""
    from modules import goods_returns, order_records, vehicle_trips
    if not isinstance(actor, dict) or actor.get("kind") not in {"loader", "dashboard"} or not actor.get("id"):
        raise DriverOrderError("Нет доступа к отправке загрузки")
    if not isinstance(operation_id, str) or not operation_id or len(operation_id) > MAX_OPERATION_ID_LENGTH:
        raise DriverOrderError("Некорректный идентификатор операции")
    with vehicle_trips._db(True), goods_returns._LOCK, order_records._LOCK, _LOCK:
        record = order_records.load_order_record(order_record_id)
        batches = order_records.loading_batches(record, route_key)
        known = next((batch for batch in batches if batch.get("operation_id") == operation_id), None)
        index = _read_index()
        if known and loading_batch_published(known):
            identity = _stable_id(known["source_assignment_id"], "loading-batch", known["id"])
            return _read(sidecar_path(identity))
        pending = batches[-1] if batches and not loading_batch_published(batches[-1]) else None
        if known and pending != known:
            raise DriverOrderConflict("Отправка уже заменена другой; обновите загрузку")
        source = _read(sidecar_path(pending["source_assignment_id"])) if pending else None
        logical_id = source.get("logical_assignment_id", source["assignment_id"]) if source else ""
        active_id = (index or {}).get("active", {}).get(logical_id)
        identity = _stable_id(source["assignment_id"], "loading-batch", pending["id"]) if pending else ""
        current = source if pending and active_id in (None, identity) else _loading_assignment(order_record_id, route_key)
        blocker = _partial_loading_blocker(current)
        if blocker:
            raise DriverOrderConflict(blocker)
        source_id = current.get("source_assignment_id") or current["assignment_id"]
        source = _read(sidecar_path(source_id))
        if source.get("order_record_id") != order_record_id or source.get("route_key") != route_key:
            raise DriverOrderConflict("Исходное назначение не соответствует маршруту")
        batch = pending or order_records.release_store_loading(
            order_record_id, revision, route_key, source_id, operation_id, actor
        )
        record = order_records.load_order_record(order_record_id)
        balance = order_records.route_return_balance(record, route_key, batch["id"])
        selected = set(batch["destination_ids"])
        snapshot = deepcopy(source)
        snapshot.update({
            "assignment_id": _stable_id(source_id, "loading-batch", batch["id"]),
            "source_assignment_id": source_id, "loading_batch_id": batch["id"],
            "batch_number": len(order_records.loading_batches(record, route_key)),
            "created_at": batch["created_at"], "revision": 0, "operations": [],
            "stores": [deepcopy(store) for store in source["stores"] if store.get("order_record_destination_id") in selected],
        })
        path = sidecar_path(snapshot["assignment_id"])
        if path.exists():
            snapshot = _read(path)
            if (snapshot.get("assignment_id") != path.stem or snapshot.get("loading_batch_id") != batch["id"]
                    or {store.get("order_record_destination_id") for store in snapshot.get("stores", [])} != selected):
                raise DriverOrderConflict("Состав сохранённой отправки изменился; восстановите назначение")
        else:
            if any(position["completed"] for destination in balance["destinations"] for position in destination["positions"]):
                raise DriverOrderConflict("Отгрузка уже начата; восстановите сохранённое назначение")
            try:
                goods_returns._snapshot_complete(snapshot, balance)
            except goods_returns.GoodsReturnError as error:
                raise DriverOrderConflict(f"Назначение не соответствует отправке: {error}") from error
            _atomic_write(path, snapshot)
        index = _read_index() or _empty_index()
        logical_id = source.get("logical_assignment_id", source_id)
        index["active"][logical_id] = snapshot["assignment_id"]
        if not any(snapshot["assignment_id"] in publication["assignment_ids"] for publication in index["publications"]):
            index["publications"].append({
                "publication_id": batch["id"], "created_at": batch["created_at"],
                "assignment_ids": [snapshot["assignment_id"]], "logical_assignment_ids": [logical_id],
            })
        _atomic_write(index_path(), index)
        return deepcopy(snapshot)


def find_assignment_for_driver(driver_id: str) -> dict | None:
    if not driver_id:
        return None
    candidates = [
        snapshot for snapshot in _snapshots(active_only=True)
        if snapshot.get("driver_id") == driver_id
    ]
    return max(candidates, key=lambda item: item.get("created_at", ""), default=None)


def completed_assignments_for_driver(driver_id: str) -> list[dict]:
    if not driver_id:
        return []
    return [
        snapshot for snapshot in _snapshots(active_only=False)
        if snapshot.get("driver_id") == driver_id
        and snapshot.get("order_record_id")
        and snapshot.get("stores")
        and all(store.get("status") == "completed" for store in snapshot["stores"])
    ]


def route_assignment_summaries(route_pairs: set[tuple[str, str]]) -> list[dict]:
    with _LOCK:
        active_ids = _active_assignment_ids()
        snapshots = _snapshots()
    result = []
    for snapshot in snapshots:
        pair = (str(snapshot.get("order_record_id", "")), str(snapshot.get("route_key", "")))
        if pair not in route_pairs:
            continue
        result.append({
            key: str(snapshot.get(key, ""))
            for key in ("order_record_id", "route_key", "assignment_id", "driver_name", "driver_id", "vehicle_name", "vehicle_plate", "created_at")
        } | {"active": active_ids is None or snapshot["assignment_id"] in active_ids})
    return sorted(result, key=lambda item: (not item["active"], item["created_at"], item["assignment_id"]))


def assignments_for_vehicle(vehicle_id: str) -> list[dict]:
    if not vehicle_id:
        return []
    return [
        snapshot for snapshot in _snapshots(active_only=True)
        if snapshot.get("vehicle_id") == vehicle_id
    ]


def assigned_vehicles() -> list[dict]:
    vehicles: dict[str, dict] = {}
    for snapshot in _snapshots(active_only=True):
        vehicle_id = str(snapshot.get("vehicle_id", ""))
        if vehicle_id:
            vehicles[vehicle_id] = {
                "id": vehicle_id,
                "name": str(snapshot.get("vehicle_name", "")),
                "plate": str(snapshot.get("vehicle_plate", "")),
            }
    return sorted(vehicles.values(), key=lambda item: (item["plate"], item["name"]))


def _find_store(snapshot: dict, store_id: str) -> dict:
    for store in snapshot.get("stores", []):
        if store.get("store_id") == store_id:
            return store
    raise DriverOrderNotFound("Магазин не найден в назначенном заказе")


def _summary(snapshot: dict) -> list[dict]:
    summary: dict[tuple[str, str], dict] = {}
    for store in snapshot.get("stores", []):
        completed = store.get("status") == "completed"
        for line in store.get("lines", []):
            name = str(line.get("name", ""))
            unit = str(line.get("unit", ""))
            item = summary.setdefault((name, unit), {
                "name": name,
                "unit": unit,
                "planned": Decimal("0"),
                "necessary": Decimal("0"),
                "completed": Decimal("0"),
            })
            planned = _decimal(line.get("required_qty", "0"))
            item["planned"] += planned
            if completed:
                item["completed"] += _decimal(line.get("actual_qty", "0"))
            else:
                item["necessary"] += planned
    return [
        {
            "name": item["name"],
            "unit": item["unit"],
            "planned": _public_number(item["planned"]),
            "necessary": _public_number(item["necessary"]),
            "completed": _public_number(item["completed"]),
        }
        for item in sorted(summary.values(), key=lambda item: (item["name"], item["unit"]))
    ]


def completed_delivery_report() -> dict:
    deliveries = []
    summary: dict[tuple[str, str], dict] = {}
    line_count = 0
    for snapshot in _snapshots(active_only=False):
        for store in snapshot.get("stores", []):
            if store.get("status") != "completed":
                continue
            lines = []
            for line in store.get("lines", []):
                try:
                    planned = _decimal(line.get("required_qty", "0"))
                    delivered = _decimal(line.get("actual_qty", "0"))
                except DriverOrderError:
                    continue
                name = str(line.get("name", ""))
                unit = str(line.get("unit", ""))
                item = summary.setdefault((name, unit), {
                    "name": name,
                    "unit": unit,
                    "planned": Decimal("0"),
                    "delivered": Decimal("0"),
                })
                item["planned"] += planned
                item["delivered"] += delivered
                lines.append({
                    "name": name,
                    "unit": unit,
                    "planned": _public_number(planned),
                    "delivered": _public_number(delivered),
                    "difference": _public_number(delivered - planned),
                })
                line_count += 1
            deliveries.append({
                "assignment_id": str(snapshot.get("assignment_id", "")),
                "publication_id": str(snapshot.get("publication_id", "")),
                "order_id": str(snapshot.get("order_id", "")),
                "route_label": str(snapshot.get("route_label", "")),
                "driver_name": str(snapshot.get("driver_name", "")),
                "vehicle_name": str(snapshot.get("vehicle_name", "")),
                "vehicle_plate": str(snapshot.get("vehicle_plate", "")),
                "store_name": str(store.get("name", "")),
                "completed_at": str(store.get("completed_at", "")),
                "lines": lines,
            })
    report_summary = [
        {
            "name": item["name"],
            "unit": item["unit"],
            "planned": _public_number(item["planned"]),
            "delivered": _public_number(item["delivered"]),
            "difference": _public_number(item["delivered"] - item["planned"]),
        }
        for item in sorted(summary.values(), key=lambda item: (item["name"], item["unit"]))
    ]
    return {
        "completed_stores": len(deliveries),
        "line_count": line_count,
        "summary": report_summary,
        "deliveries": sorted(deliveries, key=lambda item: item["completed_at"], reverse=True),
    }


def public_assignments_for_vehicle(vehicle_id: str) -> list[dict]:
    return [public_snapshot(snapshot) for snapshot in assignments_for_vehicle(vehicle_id)]


def public_snapshot(snapshot: dict) -> dict:
    payload = dict(snapshot)
    payload.pop("operations", None)
    payload["summary"] = _summary(snapshot)
    payload["completed_stores"] = sum(
        store.get("status") == "completed" for store in snapshot.get("stores", [])
    )
    payload["store_count"] = len(snapshot.get("stores", []))
    return payload


def complete_store(
    path_or_id: str | Path,
    store_id: str,
    measurements: dict[str, Any],
    operation_id: str,
    base_revision: int,
    photo: str = "",
    legacy_quantities: bool = False,
) -> dict:
    if not isinstance(operation_id, str) or not operation_id or len(operation_id) > MAX_OPERATION_ID_LENGTH:
        raise DriverOrderError("Некорректный идентификатор операции")
    if not isinstance(measurements, dict):
        raise DriverOrderError("Некорректные данные отгрузки")
    if not isinstance(base_revision, int) or base_revision < 0:
        raise DriverOrderError("Некорректная ревизия")

    path = sidecar_path(str(path_or_id)) if not isinstance(path_or_id, Path) else path_or_id
    with _LOCK:
        snapshot = _read(path)
        for operation in snapshot.get("operations", []):
            if operation.get("operation_id") == operation_id:
                return public_snapshot(snapshot)
        if snapshot.get("revision") != base_revision:
            raise DriverOrderConflict("Заказ уже изменён на другом устройстве")

        store = _find_store(snapshot, store_id)
        if store.get("status") == "completed":
            raise DriverOrderConflict("Магазин уже подтверждён")
        lines = store.get("lines", [])
        expected_ids = {line.get("line_id") for line in lines}
        if set(measurements) != expected_ids:
            raise DriverOrderError("Заполните отгрузку для каждой позиции")
        for line in lines:
            measurement = (
                _legacy_measurement(measurements[line["line_id"]])
                if legacy_quantities
                else _measurement(measurements[line["line_id"]])
            )
            line["shipment_measurement"] = measurement
            line["actual_qty"] = measurement["net_weight_kg"]
        store["status"] = "completed"
        store["completed_at"] = _now()
        store["photo"] = str(photo or "")
        snapshot["revision"] = int(snapshot.get("revision", 0)) + 1
        snapshot.setdefault("operations", []).append({
            "operation_id": operation_id,
            "store_id": store_id,
            "revision": snapshot["revision"],
            "created_at": _now(),
        })
        _atomic_write(path, snapshot)
        return public_snapshot(snapshot)


def append_store_delivery(assignment_id: str, store_id: str, measurements: dict, operation_id: str) -> dict:
    """Add a later visit to a completed shop without counting a second shop."""
    if not isinstance(operation_id, str) or not operation_id or len(operation_id) > MAX_OPERATION_ID_LENGTH:
        raise DriverOrderError("Некорректный идентификатор операции")
    path = sidecar_path(assignment_id)
    with _LOCK:
        snapshot = _read(path)
        store = _find_store(snapshot, store_id)
        expected = {line["line_id"] for line in store["lines"]}
        if store["status"] != "completed" or not isinstance(measurements, dict) or set(measurements) != expected:
            raise DriverOrderConflict("Повторная поставка не соответствует магазину")
        validated = {identity: _measurement(value) for identity, value in measurements.items()}
        if any(item.get("operation_id") == operation_id for item in snapshot.get("operations", [])):
            visits = store.get("additional_deliveries", [])
            if not any(item["operation_id"] == operation_id and item["measurements"] == validated for item in visits):
                raise DriverOrderConflict("Повторная операция содержит другие данные")
            return public_snapshot(snapshot)
        for line in store["lines"]:
            line["actual_qty"] = _quantity(_decimal(line["actual_qty"]) + _decimal(validated[line["line_id"]]["net_weight_kg"]))
        store.setdefault("additional_deliveries", []).append({"operation_id": operation_id, "at": _now(), "measurements": validated})
        snapshot["revision"] += 1
        snapshot.setdefault("operations", []).append({"operation_id": operation_id, "store_id": store_id,
                                                       "revision": snapshot["revision"], "created_at": _now()})
        _atomic_write(path, snapshot)
        return public_snapshot(snapshot)
